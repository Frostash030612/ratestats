#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate RateStats-only IA presentation PPTX (per ppt要求.docx)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "docs" / "IA_Reports"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x0B, 0x2C, 0x5E)
NAVY_LIGHT = RGBColor(0x1A, 0x4A, 0x8A)
GOLD = RGBColor(0xB8, 0x8A, 0x2E)
GRAY = RGBColor(0x44, 0x44, 0x44)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_run(run, text: str, *, size: int = 18, bold: bool = False, color=GRAY) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _add_title_bar(slide, prs, title: str, subtitle: str | None = None) -> None:
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.15))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()

    accent = slide.shapes.add_shape(1, Inches(0), Inches(1.15), prs.slide_width, Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = GOLD
    accent.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.55))
    run = box.text_frame.paragraphs[0].add_run()
    _set_run(run, title, size=26, bold=True, color=WHITE)

    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(12), Inches(0.35))
        r2 = box2.text_frame.paragraphs[0].add_run()
        _set_run(r2, subtitle, size=12, color=RGBColor(0xD0, 0xD8, 0xE8))


def _add_bullets(slide, bullets: list[str], *, left=0.6, top=1.5, width=12.0, height=5.0, size=18) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        _set_run(run, f"•  {text}", size=size, color=GRAY)


def _add_footer(slide, prs, page: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(11), Inches(0.3))
    run = box.text_frame.paragraphs[0].add_run()
    _set_run(
        run,
        "RateStats  |  Bank of China Singapore  |  Xu Wenzhe",
        size=10,
        color=RGBColor(0x88, 0x88, 0x88),
    )
    box2 = slide.shapes.add_textbox(Inches(12.0), Inches(6.9), Inches(0.8), Inches(0.3))
    p2 = box2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    _set_run(r2, page, size=10, color=RGBColor(0x88, 0x88, 0x88))


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    s = _blank(prs)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    gold = s.shapes.add_shape(1, 0, Inches(5.8), prs.slide_width, Inches(0.08))
    gold.fill.solid()
    gold.fill.fore_color.rgb = GOLD
    gold.line.fill.background()

    t = s.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _set_run(r, "RateStats", size=40, bold=True, color=WHITE)

    t2 = s.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(11.5), Inches(0.7))
    p2 = t2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    _set_run(
        r2,
        "From Manual Collection → Python Automation → AI Search",
        size=20,
        color=RGBColor(0xD0, 0xD8, 0xE8),
    )

    t3 = s.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.5), Inches(1.3))
    tf3 = t3.text_frame
    for i, line in enumerate(
        [
            "Xu Wenzhe",
            "Bank of China Singapore Branch  ·  Financial Technology Department",
            "NUS-ISS Industrial Attachment Presentation",
        ]
    ):
        p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        _set_run(run, line, size=16, color=RGBColor(0xC8, 0xD0, 0xE0))

    # 2 Agenda (aligned to ppt要求)
    s = _blank(prs)
    _add_title_bar(s, prs, "Agenda", "Aligned to presentation format · ~10 minutes")
    _add_bullets(
        s,
        [
            "Project Background",
            "My Project Role",
            "Development Phases — Requirements → Design → Implementation → Testing → Deployment",
            "Project Status",
            "Business Value & Value Added by Intern (Before / After)",
            "Demo (≤ 2 minutes, optional)",
        ],
        top=1.55,
        size=20,
    )
    _add_footer(s, prs, "2")

    # 3 Background
    s = _blank(prs)
    _add_title_bar(s, prs, "Project Background", "How the project evolved")
    _add_bullets(
        s,
        [
            "Business need: collect deposit / promo rates across ~17 Singapore banks (~47 rate nodes)",
            "Stage 0 — Fully manual: staff open bank sites, copy rates, maintain Excel by hand",
            "Stage 1 — Python app I built: fetch / parse / generate MarketRateData + Rainbow Excel",
            "          (URLs still fed manually via url_params.xlsx — fragile when banks revamp)",
            "Stage 2 — Add AI Search into that pipeline: discover URLs → validate → reuse existing parsers",
            "Constraint: controlled AI (search + scoring) — not unrestricted cloud LLM agents",
        ],
        top=1.4,
        size=16,
    )
    _add_footer(s, prs, "3")

    # 4 Role
    s = _blank(prs)
    _add_title_bar(s, prs, "My Project Role", "Built the Python product first · then added AI Search")
    _add_bullets(
        s,
        [
            "Started as Python developer: turn a manual rate-collection job into a working application",
            "Built extractors (HTML / PDF / API) and Excel outputs business users actually run",
            "Then extended the same pipeline with AI URL discovery (Vertex + provider benchmarks)",
            "Added hybrid URL selection: domain rules + ML ranker when AI candidates are ambiguous",
            "Packaged batch runs, URL health checks, scheduled output under runs/YYYYMMDD/",
            "Also supported RPA ops bots — today’s focus is RateStats only",
        ],
        top=1.35,
        size=16,
    )
    _add_footer(s, prs, "4")

    # 5 Requirements / Design
    s = _blank(prs)
    _add_title_bar(s, prs, "Development Phases — Requirements & Design", "Python foundation → AI extension")
    _add_bullets(
        s,
        [
            "Phase A (Python): automate fetch + parse + Excel — replace hand-copying of rates",
            "Phase A still needed: reliable bank parsers (HTML DOM, PDF tables, JSON APIs)",
            "Phase B (AI): stop hand-feeding every URL — search latest pages per destination",
            "Phase B also: reject wrong links (wrong bank / Premier vs board / PDF vs HTML)",
            "Final design: AI Discovery + Scoring plugged into the existing Extract & Output engine",
            "Scenario: ICBC needs SGD board/promo + FCY board/promo — same pattern × ~17 banks",
        ],
        top=1.35,
        size=16,
    )
    _add_footer(s, prs, "5")

    # 6 Architecture / Implementation
    s = _blank(prs)
    _add_title_bar(
        s,
        prs,
        "Implementation — End-to-End Architecture",
        "AI Search added on top of the Python extract/output I built first",
    )
    labels = [
        ("1. Discover", "AI Search*\ncandidate URLs"),
        ("2. Score", "Rules + ML\nhost whitelist"),
        ("3. Extract*", "HTML / PDF / API\nPython parsers"),
        ("4. Output*", "MarketRateData\n+ Rainbow Excel"),
    ]
    box_w = Inches(2.6)
    gap = Inches(0.35)
    start_x = Inches(0.9)
    y = Inches(1.9)
    for i, (h, body) in enumerate(labels):
        x = start_x + i * (box_w + gap)
        shape = s.shapes.add_shape(1, x, y, box_w, Inches(2.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT_BG
        shape.line.color.rgb = NAVY
        tb = s.shapes.add_textbox(x + Inches(0.15), y + Inches(0.3), box_w - Inches(0.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        _set_run(r, h, size=18, bold=True, color=NAVY)
        for line in body.split("\n"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            _set_run(r2, line, size=14, color=GRAY)
        if i < 3:
            ax = x + box_w + Inches(0.02)
            at = s.shapes.add_textbox(ax, y + Inches(0.9), gap, Inches(0.5))
            ap = at.text_frame.paragraphs[0]
            ap.alignment = PP_ALIGN.CENTER
            ar = ap.add_run()
            _set_run(ar, "→", size=22, bold=True, color=GOLD)

    note = s.shapes.add_textbox(Inches(0.9), Inches(4.5), Inches(11.5), Inches(1.5))
    tf = note.text_frame
    tf.word_wrap = True
    lines = [
        "* Steps 3–4 existed first as a Python application; Steps 1–2 were added later",
        "Key modules: bank_fetch_and_extract.py · bank_extractors/ · then vertex_url_discovery.py",
        "url_host_rules.py · url_discovery_pick.py · RateStats_ML/picker.py",
        "Challenge: AI returns HLF for ICBC → domain whitelist before ML",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, f"•  {line}", size=15, color=GRAY)
    _add_footer(s, prs, "6")

    # 7 Testing & Deployment
    s = _blank(prs)
    _add_title_bar(s, prs, "Testing & Deployment", "How we know it works · how it runs")
    _add_bullets(
        s,
        [
            "Testing: compare AI-picked URLs vs manual golden urls in url_params.xlsx",
            "Testing: picker_eval reports · multi-provider benchmark (Vertex / Serper / Brave / Tavily)",
            "Testing: URL health checks (status / redirect) before scheduled runs",
            "Deployment: batch scripts + Windows Task Scheduler → dated runs/YYYYMMDD/",
            "Deliverables: MarketRateData_*.xlsx + Rainbow Table Excel for business users",
        ],
        top=1.4,
        size=17,
    )
    _add_footer(s, prs, "7")

    # 8 Project Status
    s = _blank(prs)
    _add_title_bar(s, prs, "Project Status", "C = Completed · I = In progress")
    # Simple status table as bullets (clear for 10-min talk)
    _add_bullets(
        s,
        [
            "Python extraction (HTML / PDF / API) — Production (C)  [built first]",
            "Batch pipeline + Excel (MarketRateData / Rainbow) — Production (C)",
            "AI URL Discovery (Vertex) — Production (C)  [added into the flow]",
            "Rule-based scoring & host validation — Production (C)",
            "Multi-AI benchmarking module — Production (C)",
            "ML hybrid URL ranker — Integration done; production tuning ongoing (I)",
        ],
        top=1.4,
        size=17,
    )
    _add_footer(s, prs, "8")

    # 9 Business Value Before/After
    s = _blank(prs)
    _add_title_bar(s, prs, "Business Value — Before vs After", "Emphasize tangible outcomes")
    # Two columns
    left = s.shapes.add_shape(1, Inches(0.6), Inches(1.6), Inches(5.7), Inches(4.5))
    left.fill.solid()
    left.fill.fore_color.rgb = LIGHT_BG
    left.line.color.rgb = RGBColor(0xAA, 0x55, 0x55)
    right = s.shapes.add_shape(1, Inches(6.9), Inches(1.6), Inches(5.7), Inches(4.5))
    right.fill.solid()
    right.fill.fore_color.rgb = LIGHT_BG
    right.line.color.rgb = RGBColor(0x2E, 0x7D, 0x4F)

    lb = s.shapes.add_textbox(Inches(0.85), Inches(1.8), Inches(5.2), Inches(4.0))
    tf = lb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "BEFORE (Manual)", size=18, bold=True, color=RGBColor(0xAA, 0x55, 0x55))
    for line in [
        "Fully manual: open sites, copy rates",
        "Heavy manpower on ~47 rate nodes",
        "Easy to miss updates / typos",
        "No reusable application pipeline",
    ]:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(10)
        r2 = p2.add_run()
        _set_run(r2, f"•  {line}", size=15, color=GRAY)

    rb = s.shapes.add_textbox(Inches(7.15), Inches(1.8), Inches(5.2), Inches(4.0))
    tf = rb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "AFTER (Python + AI)", size=18, bold=True, color=RGBColor(0x2E, 0x7D, 0x4F))
    for line in [
        "Python app: fetch / parse → Excel",
        "AI Search finds URLs; rules/ML check",
        "Scheduled runs; fewer missing rows",
        "~95% URL hit; Vertex chosen by data",
    ]:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(10)
        r2 = p2.add_run()
        _set_run(r2, f"•  {line}", size=15, color=GRAY)
    _add_footer(s, prs, "9")

    # 10 Value added by intern
    s = _blank(prs)
    _add_title_bar(s, prs, "Value Added by the Intern", "Beyond coding tickets")
    _add_bullets(
        s,
        [
            "Delivered a usable Python product first — not only a research prototype",
            "Grew it from scripts to a schedulable application business users can run",
            "Then plugged AI Search into that real pipeline (not a separate demo)",
            "Designed hybrid rules + ML when pure AI candidates were wrong",
            "Benchmarking gave evidence to choose Vertex; docs enable handover",
        ],
        top=1.45,
        size=17,
    )
    _add_footer(s, prs, "10")

    # 11 Demo
    s = _blank(prs)
    _add_title_bar(s, prs, "Demo (≤ 2 minutes)", "Scenario-focused · skip login / CRUD noise")
    ph = s.shapes.add_shape(1, Inches(1.5), Inches(1.7), Inches(10.3), Inches(4.3))
    ph.fill.solid()
    ph.fill.fore_color.rgb = LIGHT_BG
    ph.line.color.rgb = NAVY_LIGHT
    box = s.shapes.add_textbox(Inches(1.8), Inches(2.8), Inches(9.7), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _set_run(r, "[ Live demo or screen-grab GIF ]", size=22, bold=True, color=NAVY)
    for line in [
        "1) Run batch discovery / fetch",
        "2) Show selected URLs vs wrong candidates (scenario)",
        "3) Open MarketRateData / Rainbow Excel output",
    ]:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(8)
        r2 = p2.add_run()
        _set_run(r2, line, size=15, color=GRAY)
    _add_footer(s, prs, "11")

    # 12 Closing
    s = _blank(prs)
    _add_title_bar(s, prs, "Conclusion & Q&A", "")
    _add_bullets(
        s,
        [
            "Journey: fully manual → Python automation (extract/Excel) → AI Search in the same flow",
            "AI is an extension of a working application — not the starting point",
            "Hybrid Rules + ML keeps URL accuracy high (~95%) under bank security constraints",
            "Business impact: less manpower, more reliable market-rate Excel — thank you / Q&A",
        ],
        top=1.65,
        size=18,
    )
    _add_footer(s, prs, "12")

    path = OUT / "IA_Presentation_RateStats_Xu_Wenzhe.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build())
