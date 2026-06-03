#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《AI 搜索全流程说明（含目录与脚本对照）》Word 文档。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:
    raise SystemExit("请先安装: pip install python-docx")

_THIS = Path(__file__).resolve().parent
_ROOT = _THIS.parent
_OUT = _THIS / "docs" / "AI搜索全流程说明_含脚本对照.docx"
_OUT_EN = _THIS / "docs" / "AI_Search_Workflow_Scripts_Guide.docx"


def _h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(10.5)


def _bullet(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _mono(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            t.rows[i + 1].cells[j].text = str(cell)


def build() -> Document:
    doc = Document()
    t = doc.add_heading("RateStats：AI 搜索全流程与脚本对照手册", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p(doc, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    _p(doc, "项目根目录：RateStats（含 RateStats_Portable、AI_Compare）")

    # ========== 目录树 ==========
    _h(doc, "一、项目目录结构（与脚本相关部分）", 1)
    _p(doc, "以下为实际参与「手动 / AI 搜索 / 抓取 / 对比」流程的目录树（省略 temp、历史输出文件夹）：")
    tree = r"""
RateStats/
├── RateStats_Portable/                    # 主生产目录
│   ├── 01_首次安装依赖.bat
│   ├── 02_一键生成Excel.bat
│   ├── 03_一键生成Market+彩虹表.bat         # 手动：Market + 彩虹表
│   ├── 04_一键巡检URL.bat
│   ├── 05_run_ai_search.bat                 # 仅 Vertex AI
│   ├── 05_一键AI搜索生成Market.bat
│   ├── 06_一键生成Market彩虹表与AI搜索.bat  # 手动 + Vertex AI
│   ├── assets/
│   │   ├── url_params.xlsx                  # 手动链接配置（人工维护）
│   │   ├── url_params.json
│   │   ├── url_params_ai.xlsx               # Vertex 单跑 AI 写入
│   │   ├── url_params_ai.json
│   │   ├── ratestatsearch-*.json            # Vertex 服务账号
│   │   └── MarketRateData_template.xlsx
│   ├── YYYYMMDD/                            # 当日输出目录
│   │   ├── MarketRateData_*_manual*.xlsx
│   │   ├── MarketRateData_*_Vertex*.xlsx
│   │   ├── MarketRateData_*_AISearch.xlsx
│   │   └── url_YYYYMMDD.xlsx
│   ├── sync_url_params_to_json.py           # 手动 xlsx → json
│   ├── sync_url_params_ai_to_json.py        # AI xlsx → json
│   ├── market_rate_data_generator.py        # 手动入口（→ bank_all_promo_rates）
│   ├── run_market_rate_ai_search.py         # Vertex：发现 + 抓取 AISearch
│   ├── vertex_url_discovery.py              # Vertex URL 发现
│   ├── vertex_search_client.py              # Vertex API 客户端
│   ├── url_discovery_pick.py                # 四家共用：打分 + 选链
│   ├── url_host_rules.py / url_path_rules.py / url_dest_intent.py
│   ├── url_pick_refinement.py / url_fallback_resolver.py
│   ├── url_discovery_config.py / url_config_loader.py
│   ├── url_key_aliases.py / url_sources.py / url_defaults.py
│   ├── bank_all_promo_rates.py              # 抓取主程序 CLI
│   ├── bank_fetch_and_extract.py            # 各银行抓取汇总
│   ├── bank_extractors/impl.py              # 各银行解析实现
│   ├── bank_output_stage.py                 # 写 Excel + 导出 url 快照
│   ├── bank_excel_data_builder.py           # payload → DataFrame
│   ├── bank_excel_template_writer.py
│   ├── export_market_url_snapshot.py
│   ├── generate_rainbow_from_market.py
│   └── test_vertex_search.py
│
└── AI_Compare/                              # 多 AI 对比（独立）
    ├── keys.txt / keys.txt.example          # Serper/Brave/Tavily Key
    ├── load_keys.py
    ├── _discovery_common.py                 # 共享发现 + QUERY + IO
    ├── run_all_providers_market.py          # 四家发现 + 抓取
    ├── run_all_providers_market.bat
    ├── compare_market_data.py               # 单家 AI vs 手动 数据对比
    ├── compare_all_providers_market.py      # 四家批量对比
    ├── compare_manual_vs_ai.py              # url_params 配置对比
    ├── run_eval.py / run_eval.bat           # 仅 URL 命中率评测
    ├── ground_truth.py / eval_metrics.py / queries.py
    ├── audit_ai_issues.py
    ├── output/<tag>/                        # 各 AI 的 url_params_ai_*.xlsx
    ├── results/                             # 对比报告 Excel
    ├── vertex/                              # Vertex 冒烟（非完整发现）
    │   ├── vertex_search_client.py
    │   └── test_vertex_search.py
    ├── serper/
    │   ├── serper_search_client.py
    │   ├── serper_url_discovery.py          # 可单独跑（走 _discovery_common）
    │   └── test_serper_search.py
    ├── brave/
    │   ├── brave_search_client.py
    │   ├── brave_url_discovery.py
    │   └── test_brave_search.py
    └── tavily/
        ├── tavily_search_client.py
        ├── tavily_url_discovery.py
        └── test_tavily_search.py
"""
    for line in tree.strip().split("\n"):
        _mono(doc, line)

    # ========== 手动流程 ==========
    _h(doc, "二、手动查询流程：逐步脚本对照", 1)
    _table(
        doc,
        ["步骤", "说明", "入口脚本（批处理/命令）", "核心 Python 模块", "输入", "输出"],
        [
            [
                "0 准备",
                "同步手动配置",
                "03_一键生成Market+彩虹表.bat（开头）",
                "sync_url_params_to_json.py\nurl_config_loader.py\nurl_key_aliases.py",
                "assets/url_params.xlsx",
                "assets/url_params.json",
            ],
            [
                "1 抓取",
                "访问各银行页/API 解析利率",
                "market_rate_data_generator.py\n（或 03 批处理）",
                "bank_all_promo_rates.py\n→ bank_cli_parser.py\n→ bank_cli_url_config_applier.py\n→ bank_cli_fetch_params.py\n→ bank_fetch_and_extract.py\n→ bank_extractors/impl.py",
                "url_params.xlsx（默认）",
                "内存 payload（各银行块）",
            ],
            [
                "2 写 Excel",
                "生成 Market 多 sheet",
                "（同上，一步完成）",
                "bank_output_stage.py\n→ bank_excel_data_builder.py\n→ bank_excel_template_writer.py\n→ bank_excel_style_helpers.py\n→ bank_xlsx_out_resolver.py",
                "payload",
                "MarketRateData_*.xlsx",
            ],
            [
                "3 URL 快照",
                "从元数据导出黄金 URL",
                "（自动）",
                "export_market_url_snapshot.py",
                "MarketRateData 元数据 sheet",
                "url_YYYYMMDD.xlsx",
            ],
            [
                "4 彩虹表",
                "可选",
                "03 批处理 第 2 步",
                "generate_rainbow_from_market.py",
                "MarketRateData_*.xlsx",
                "彩虹表_*.xlsx",
            ],
            [
                "5 邮件",
                "可选",
                "03/06 批处理",
                "send_report_email.py",
                "Market + 彩虹表",
                "发送报告",
            ],
        ],
    )
    _p(doc, "手动链路辅助/配置模块：url_sources.py（默认 URL）、url_defaults.py、url_args_override.py、bank_constants.py、bank_safety.py")

    # ========== AI 发现 ==========
    _h(doc, "三、AI 搜索发现流程：逐步脚本对照", 1)

    _h(doc, "3.1 公共准备（四家 AI 均会用到）", 2)
    _table(
        doc,
        ["模块", "作用"],
        [
            ["AI_Compare/load_keys.py", "从 keys.txt 加载 Serper/Brave/Tavily API Key"],
            ["AI_Compare/_discovery_common.py", "discover_all 循环、QUERY_BY_DEST、写 xlsx 报表"],
            ["RateStats_Portable/url_discovery_pick.py", "score_url + pick_best_url（选链核心）"],
            ["RateStats_Portable/url_host_rules.py", "银行域名约束"],
            ["RateStats_Portable/url_path_rules.py", "路径加减分"],
            ["RateStats_Portable/url_dest_intent.py", "意图软加分 / 硬拒绝"],
            ["RateStats_Portable/url_pick_refinement.py", "宽松/严格模式下的候选精炼"],
            ["RateStats_Portable/url_fallback_resolver.py", "fallback 优先 url_params.xlsx"],
            ["RateStats_Portable/url_discovery_config.py", "relaxed / strict 模式开关"],
            ["RateStats_Portable/url_key_aliases.py", "dest ↔ 中文 key 映射"],
        ],
    )

    _h(doc, "3.2 Vertex AI Search", 2)
    _table(
        doc,
        ["步骤", "入口脚本", "专用模块", "输出"],
        [
            [
                "发现",
                "run_market_rate_ai_search.py（--discover-only 可只发现）\n或 run_all_providers_market.py --providers vertex\n或 05/06 批处理",
                "vertex_url_discovery.py\n→ vertex_search_client.py（search_url, extract_links）\n→ url_discovery_pick.pick_best_url",
                "assets/url_params_ai.xlsx\nassets/ai_search_discovered_*.xlsx\n或 AI_Compare/output/<tag>/url_params_ai_vertex.xlsx",
            ],
            [
                "同步 JSON",
                "（发现后自动）",
                "sync_url_params_ai_to_json.py",
                "assets/url_params_ai.json",
            ],
        ],
    )

    _h(doc, "3.3 Serper / Brave / Tavily", 2)
    _table(
        doc,
        ["Provider", "搜索客户端", "发现入口（可单独跑）", "批量入口"],
        [
            [
                "Serper",
                "serper/serper_search_client.py",
                "serper/serper_url_discovery.py",
                "run_all_providers_market.py → _discovery_common.discover_all(search_fn=search_urls)",
            ],
            [
                "Brave",
                "brave/brave_search_client.py",
                "brave/brave_url_discovery.py",
                "同上",
            ],
            [
                "Tavily",
                "tavily/tavily_search_client.py",
                "tavily/tavily_url_discovery.py",
                "同上",
            ],
        ],
    )
    _p(doc, "四家批量发现时，配置与发现报告写入：AI_Compare/output/<RUN_TAG>/url_params_ai_<provider>.xlsx")

    # ========== AI 抓取 ==========
    _h(doc, "四、AI 抓取生成 Market：逐步脚本对照", 1)
    _table(
        doc,
        ["步骤", "入口", "脚本链", "输入配置", "输出"],
        [
            [
                "Vertex 单跑",
                "run_market_rate_ai_search.py（无 --discover-only）\n05_run_ai_search.bat",
                "sync_url_params_ai_to_json.py\n→ bank_all_promo_rates.py（--url-config url_params_ai.xlsx）\n→ 同手动抓取链",
                "assets/url_params_ai.xlsx",
                "MarketRateData_*_AISearch.xlsx",
            ],
            [
                "四家批量",
                "run_all_providers_market.py",
                "run_one_provider → sync_url_json\n→ fetch_market → bank_all_promo_rates.py",
                "output/<tag>/url_params_ai_<provider>.xlsx",
                "MarketRateData_<tag>_Vertex|Serper|Brave|Tavily.xlsx",
            ],
        ],
    )

    # ========== 对比 ==========
    _h(doc, "五、对比与评测：逐步脚本对照", 1)
    _table(
        doc,
        ["目的", "脚本", "说明", "输出目录"],
        [
            [
                "只比 URL 命中率",
                "run_eval.py\nrun_compare_all.bat",
                "ground_truth.py + eval_metrics.py + providers/*.py",
                "AI_Compare/results/ai_provider_compare_*.xlsx",
            ],
            [
                "手动 vs 单家 AI 数据",
                "compare_market_data.py",
                "对比三主表行 + 利率 + 元数据 URL",
                "results/.../market_data_compare_*.xlsx",
            ],
            [
                "手动 vs 四家 AI",
                "compare_all_providers_market.py",
                "汇总 summary + pivot",
                "results/.../market_compare_all_providers_*.xlsx",
            ],
            [
                "url_params 配置差异",
                "compare_manual_vs_ai.py",
                "手动 xlsx vs url_params_ai",
                "manual_vs_ai_urls_*.xlsx",
            ],
            [
                "问题审计",
                "audit_ai_issues.py",
                "链接问题 + 缺行归因",
                "ai_issues_audit_*.xlsx",
            ],
        ],
    )

    # ========== 端到端 ==========
    _h(doc, "六、端到端推荐命令（含脚本顺序）", 1)
    _p(doc, "A. 手动 + 四家 AI + 与手动数据对比（完整评测）")
    steps_a = [
        "cd RateStats_Portable",
        "python sync_url_params_to_json.py",
        "python market_rate_data_generator.py --xlsx-out .\\YYYYMMDD\\MarketRateData_<tag>_manual.xlsx",
        "cd ..\\AI_Compare",
        "python run_all_providers_market.py --run-tag <tag> --out-dir ..\\RateStats_Portable\\YYYYMMDD --skip-missing-keys --intent-mode relaxed",
        "python compare_all_providers_market.py --manual ..\\...\\MarketRateData_<tag>_manual.xlsx --ai-dir ..\\...\\YYYYMMDD --ai-glob MarketRateData_<tag>* --exclude-name-substr manual",
    ]
    for s in steps_a:
        _mono(doc, s)

    _p(doc, "B. 仅 Vertex（Portable 生产）")
    steps_b = [
        "cd RateStats_Portable",
        "python run_market_rate_ai_search.py --intent-mode relaxed",
        "或：06_一键生成Market彩虹表与AI搜索.bat（含手动+Vertex）",
    ]
    for s in steps_b:
        _mono(doc, s)

    # ========== 依赖关系简图 ==========
    _h(doc, "七、核心调用关系（简图）", 1)
    _mono(
        doc,
        """
[手动]
  03.bat / market_rate_data_generator.py
    → bank_all_promo_rates.main()
      → apply_url_config (url_params.xlsx)
      → fetch_and_extract (bank_fetch_and_extract + impl)
      → run_output_stage → write_rates_excel
      → export_market_url_snapshot

[AI 发现 - Vertex]
  run_market_rate_ai_search / run_all_providers_market
    → vertex_url_discovery.discover_all()
      → vertex_search_client.search_url
      → url_discovery_pick.pick_best_url
      → write_url_params_xlsx

[AI 发现 - Serper/Brave/Tavily]
  run_all_providers_market
    → _discovery_common.discover_all(search_fn=xxx_search_client.search_urls)
      → url_discovery_pick.pick_best_url (同上)

[AI 抓取]
  bank_all_promo_rates.main(--url-config url_params_ai_*.xlsx)
    → （与手动相同抓取链）

[对比]
  compare_market_data / compare_all_providers_market
    → export_market_url_snapshot.dest_urls_from_market_meta (读元数据)
""".strip(),
    )

    _h(doc, "八、配置文件与数据文件速查", 1)
    _table(
        doc,
        ["文件", "谁维护", "用途"],
        [
            ["assets/url_params.xlsx", "人工", "手动抓取唯一配置源"],
            ["assets/url_params_ai.xlsx", "Vertex 单跑生成", "Portable AI 抓取用"],
            ["AI_Compare/output/<tag>/url_params_ai_*.xlsx", "各 AI 发现生成", "四家批量抓取用"],
            ["AI_Compare/keys.txt", "人工", "Serper/Brave/Tavily Key"],
            ["assets/ratestatsearch-*.json", "GCP 服务账号", "Vertex 认证"],
        ],
    )

    _h(doc, "九、批处理文件速查", 1)
    _table(
        doc,
        ["批处理", "目录", "等价命令行主脚本"],
        [
            ["03_一键生成Market+彩虹表.bat", "Portable", "sync + market_rate_data_generator + generate_rainbow"],
            ["05_run_ai_search.bat", "Portable", "run_market_rate_ai_search.py"],
            ["06_一键生成Market彩虹表与AI搜索.bat", "Portable", "03 全流程 + run_market_rate_ai_search"],
            ["run_all_providers_market.bat", "AI_Compare", "run_all_providers_market.py"],
            ["run_eval.bat / run_compare_all.bat", "AI_Compare", "run_eval.py"],
        ],
    )

    doc.add_page_break()
    _p(doc, "附录：重新生成本文档请运行 AI_Compare/generate_ai_workflow_doc.py")
    return doc


def main() -> int:
    _OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build()
    doc.save(_OUT)
    doc.save(_OUT_EN)
    print(f"[OK] {_OUT}")
    print(f"[OK] {_OUT_EN}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
