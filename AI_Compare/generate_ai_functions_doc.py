#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《AI 脚本函数说明手册》Word 文档。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:
    raise SystemExit("请先安装: pip install python-docx")

_THIS = Path(__file__).resolve().parent
_OUT = _THIS / "docs" / "AI脚本函数说明手册.docx"
_OUT_EN = _THIS / "docs" / "AI_Functions_Reference.docx"

# (章节标题, 文件名说明, [(函数名, 作用), ...])
_SECTIONS: list[tuple[str, str, list[tuple[str, str]]]] = [
    (
        "一、RateStats_Portable — 发现与选链",
        "",
        [],
    ),
    (
        "url_discovery_config.py — 全局模式",
        "RateStats_Portable/url_discovery_config.py",
        [
            ("get_intent_mode()", "读环境变量 RATESTATS_INTENT_MODE，返回 relaxed 或 strict"),
            ("is_relaxed_mode()", "是否为宽松模式（默认，利于银行改版）"),
        ],
    ),
    (
        "url_discovery_pick.py — 选链核心（四家 AI 共用）",
        "RateStats_Portable/url_discovery_pick.py",
        [
            ("score_url(url, dest, reference_url=...)", "对候选 URL 综合打分（域名、路径、促销/挂牌词、ICBC 页 ID、意图软加分等）"),
            ("pick_best_url(candidates, dest, fallback)", "从搜索返回链接中选最终 URL；返回 (url, source_tag)"),
        ],
    ),
    (
        "url_host_rules.py — 银行域名",
        "RateStats_Portable/url_host_rules.py",
        [
            ("host_fragments_for_dest(dest)", "该 dest 允许的域名片段"),
            ("url_matches_dest(dest, url)", "URL 是否属于目标银行（含 BOC/SCB 的 /sg/ 校验）"),
            ("filter_urls_for_dest(dest, urls)", "过滤出合法银行域名的 URL 列表"),
            ("assert_all_dests_have_host_rules()", "启动校验：每个 dest 必须有域名规则"),
        ],
    ),
    (
        "url_path_rules.py — 路径加减分",
        "RateStats_Portable/url_path_rules.py",
        [
            ("path_score_adjustment(dest, url)", "按银行/产品页路径加减分（Maybank JSP、OCBC 个人页、DBS API 等）"),
        ],
    ),
    (
        "url_dest_intent.py — 路径意图",
        "RateStats_Portable/url_dest_intent.py",
        [
            ("_norm(url)", "URL 转小写"),
            ("url_meets_dest_intent(dest, url)", "strict：路径必须满足意图；relaxed：仅排除硬拒绝"),
            ("filter_by_intent(dest, urls)", "过滤满足意图的 URL"),
            ("_path_tokens(url)", "提取路径 token，用于与手动 URL 比对相似度"),
            ("url_hard_reject(dest, url)", "宽松模式下仍一票否决的错页"),
            ("intent_score_adjustment(dest, url, reference_url=...)", "宽松模式软加分/减分（参考 fallback URL）"),
        ],
    ),
    (
        "url_pick_refinement.py — 候选精炼",
        "RateStats_Portable/url_pick_refinement.py",
        [
            ("refine_ranked_candidates_strict(...)", "严格：过滤不符合意图；特殊 dest 强制 API/JSP/zh-sg 等"),
            ("refine_ranked_candidates_relaxed(...)", "宽松：仅去掉硬拒绝项"),
            ("refine_ranked_candidates(...)", "按当前模式分发到 strict/relaxed"),
        ],
    ),
    (
        "url_fallback_resolver.py — 回退 URL",
        "RateStats_Portable/url_fallback_resolver.py",
        [
            ("_code_defaults()", "代码内 DEFAULT_* URL 表"),
            ("build_fallback_map(manual_xlsx=...)", "优先 url_params.xlsx，未配置 dest 用代码默认"),
        ],
    ),
    (
        "vertex_search_client.py — Vertex API",
        "RateStats_Portable/vertex_search_client.py",
        [
            ("_key_file()", "服务账号 JSON 路径"),
            ("get_access_token()", "OAuth2 access token"),
            ("search_url(query, page_size=8)", "调用 Discovery Engine 搜索，返回 JSON"),
            ("extract_links(payload)", "从 JSON 响应解析 http 链接"),
        ],
    ),
    (
        "vertex_url_discovery.py — Vertex 全量发现",
        "RateStats_Portable/vertex_url_discovery.py",
        [
            ("DiscoverResult", "单条发现结果 dataclass（key/dest/query/url/source/candidates）"),
            ("_canonical_friendly_keys()", "dest → Excel 列名 key"),
            ("pick_best_url(...)", "封装共享选链，source ai 重命名为 vertex"),
            ("discover_all(...)", "遍历全部 dest：Vertex 搜索 → 选链"),
            ("write_url_params_xlsx(...)", "写 url_params_ai.xlsx"),
            ("write_discovery_report(...)", "写发现明细报告"),
            ("run_discovery(...)", "一键发现 + 写配置与报告"),
        ],
    ),
    (
        "run_market_rate_ai_search.py — Vertex 生产入口",
        "RateStats_Portable/run_market_rate_ai_search.py",
        [
            ("_sync_ai_json()", "url_params_ai.xlsx → url_params_ai.json"),
            ("_run_fetch(market_out)", "用 AI 配置调用 bank_all_promo_rates 抓取 Market"),
            ("main()", "CLI：发现（可选）→ 同步 JSON → 抓取"),
        ],
    ),
    (
        "sync_url_params_ai_to_json.py",
        "RateStats_Portable/sync_url_params_ai_to_json.py",
        [
            ("main()", "assets/url_params_ai.xlsx → assets/url_params_ai.json"),
        ],
    ),
    (
        "二、AI_Compare — 四家对比与发现",
        "",
        [],
    ),
    (
        "_discovery_common.py — 共享发现",
        "AI_Compare/_discovery_common.py",
        [
            ("DiscoverResult", "发现结果结构（与 Vertex 版字段一致）"),
            ("canonical_friendly_keys()", "dest → Excel 列名 key"),
            ("discover_all(search_fn, source_tag, ...)", "对每个 dest 调 search_fn(query) 再 pick_best_url"),
            ("write_url_params_xlsx(...)", "写 AI 配置 xlsx（key + url）"),
            ("write_discovery_report(...)", "写发现明细报表"),
            ("default_output_paths(provider_dir)", "返回默认 url_params_ai.xlsx 与发现报告路径"),
            ("run_provider_discovery(...)", "发现 + 写两份 xlsx（单 provider 目录入口）"),
        ],
    ),
    (
        "load_keys.py",
        "AI_Compare/load_keys.py",
        [
            ("_parse_line(line)", "解析 keys.txt 中 KEY=VALUE 行"),
            ("load_keys_from_file(path)", "加载单个文件到环境变量"),
            ("ensure_keys_loaded(...)", "自动查找 keys.txt / keys.env"),
            ("main()", "命令行检查 Key 是否已加载"),
        ],
    ),
    (
        "run_all_providers_market.py — 四家批量",
        "AI_Compare/run_all_providers_market.py",
        [
            ("_provider_configured(name)", "检查 Vertex 服务账号或 Serper/Brave/Tavily Key"),
            ("_discover_api_provider(...)", "Serper/Brave/Tavily 发现并写 xlsx"),
            ("sync_url_json(xlsx, json)", "单 provider 配置 xlsx → json"),
            ("fetch_market(url_xlsx, market_out)", "用指定 AI 配置抓取 Market Excel"),
            ("run_one_provider(...)", "单家全流程：发现 → 同步 → 抓取"),
            ("main()", "CLI 按 --providers 循环四家"),
        ],
    ),
    (
        "serper / brave / tavily — 搜索客户端（结构相同）",
        "AI_Compare/<provider>/<provider>_search_client.py",
        [
            ("_api_key()", "从环境变量读取 API Key"),
            ("search(query, page_size)", "HTTP 搜索，返回原始 JSON"),
            ("extract_links(payload)", "从响应解析链接列表"),
            ("search_urls(query)", "发现流程入口：搜索并直接返回 URL 列表"),
        ],
    ),
    (
        "serper / brave / tavily — 单独发现脚本",
        "AI_Compare/<provider>/<provider>_url_discovery.py",
        [
            ("main()", "仅跑该 provider 全量发现，输出到 <provider>/output/"),
        ],
    ),
    (
        "三、对比与评测脚本",
        "",
        [],
    ),
    (
        "compare_market_data.py — 手动 vs AI 数据",
        "AI_Compare/compare_market_data.py",
        [
            ("_pct_cols(df)", "找出利率百分比列（列名以 _pct 结尾）"),
            ("_row_key(row, key_cols)", "用非利率列拼成行唯一键，用于行对齐"),
            ("_norm_rate(v) / _rates_equal(a, b, tol)", "利率规范化与容差内相等判断"),
            ("compare_sheet(...)", "单 sheet 行级对比，返回明细与汇总"),
            ("compare_metadata_urls(manual, ai)", "对比两份 Market 元数据中的抓取 URL"),
            ("run_compare(...)", "生成 market_data_compare_*.xlsx"),
            ("main()", "CLI：--manual 与 --ai"),
        ],
    ),
    (
        "compare_all_providers_market.py",
        "AI_Compare/compare_all_providers_market.py",
        [
            ("_provider_label(ai_path)", "从 Market 文件名推断 provider 显示名"),
            ("main()", "批量对多家 AI 跑 run_compare 并汇总总表"),
        ],
    ),
    (
        "compare_manual_vs_ai.py",
        "AI_Compare/compare_manual_vs_ai.py",
        [
            ("run(manual_xlsx, ai_xlsx, out_dir)", "逐 dest 对比手动快照与 AI 配置 URL"),
            ("main()", "CLI"),
        ],
    ),
    (
        "audit_ai_issues.py",
        "AI_Compare/audit_ai_issues.py",
        [
            ("audit_urls(...)", "链接为空、错域名、错行、同银行不同页"),
            ("audit_data_gaps(...)", "缺银行整行、各银行行数差异"),
            ("_meta_errors(xlsx)", "元数据 sheet 中的抓取失败项"),
            ("main()", "写 ai_issues_audit_*.xlsx"),
        ],
    ),
    (
        "ground_truth.py — 评测黄金答案",
        "AI_Compare/ground_truth.py",
        [
            ("load_gold(path)", "从 xlsx 或 json 读取 dest→url"),
            ("find_latest_url_snapshot(...)", "查找最新 url_YYYYMMDD.xlsx"),
            ("resolve_gold_path(...)", "解析 run_eval 使用的黄金答案路径"),
        ],
    ),
    (
        "eval_metrics.py",
        "AI_Compare/eval_metrics.py",
        [
            ("normalize_url(u)", "URL 规范化（scheme/host/path）"),
            ("rank_in_topn(gold, candidates)", "黄金 URL 在候选中的 1-based 排名，未命中为 0"),
            ("hit_at_k(rank, k)", "Hit@K 指标"),
            ("mrr(rank)", "Mean Reciprocal Rank"),
            ("same_host_rank(gold, candidates)", "同域名命中排名（不要求路径相同）"),
        ],
    ),
    (
        "run_eval.py — 仅 URL 命中率",
        "AI_Compare/run_eval.py",
        [
            ("_build_providers(names)", "实例化已配置 Key 的 provider"),
            ("_effective_sleep(...)", "Brave 限流时自动抬高请求间隔"),
            ("_cache_raw(...)", "缓存单次搜索原始 JSON"),
            ("run(providers, topn, ...)", "各 dest 搜索 vs 黄金 URL，写 ai_provider_compare_*.xlsx"),
            ("cmd_list()", "打印各 provider Key 配置状态"),
            ("main()", "CLI：--list / --auto / --providers"),
        ],
    ),
]


def _h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(10.5)


def _mono(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            t.rows[i + 1].cells[j].text = str(cell)


def build() -> Document:
    doc = Document()
    title = doc.add_heading("RateStats：AI 脚本函数说明手册", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p(doc, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    _p(
        doc,
        "说明：各函数已在对应 .py 源码中补充中文 docstring；本文档为 Word 速查版。"
        " 项目路径：RateStats/RateStats_Portable 与 RateStats/AI_Compare。",
    )

    _h(doc, "调用关系（发现 → 抓取 → 对比）", 1)
    flow = """
search_urls / search_url + extract_links
        ↓
discover_all（vertex_url_discovery 或 _discovery_common）
        ↓
pick_best_url → url_discovery_pick
        ├─ score_url ← url_host_rules, url_path_rules, url_dest_intent
        └─ refine_ranked_candidates ← url_pick_refinement
        ↓
write_url_params_xlsx → bank_all_promo_rates（--url-config）
        ↓
compare_market_data / compare_all_providers_market
""".strip()
    _mono(doc, flow)

    for heading, filepath, rows in _SECTIONS:
        if not rows and not filepath:
            _h(doc, heading, 1)
            continue
        _h(doc, heading, 2)
        if filepath:
            _p(doc, f"文件：{filepath}")
        _table(doc, ["函数 / 类", "作用"], [[a, b] for a, b in rows])

    _h(doc, "附录：查看源码注释", 1)
    _p(doc, "在 IDE 中打开上述 .py 文件，悬停函数名可查看 docstring。")
    _p(doc, "重新生成本 Word：在 AI_Compare 目录执行")
    _mono(doc, "python generate_ai_functions_doc.py")
    _p(doc, "Markdown 版：AI_Compare/docs/AI函数说明.md")
    _p(doc, "全流程 Word：AI_Compare/docs/AI搜索全流程说明_含脚本对照.docx")

    return doc


def main() -> int:
    _OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build()
    doc.save(_OUT)
    doc.save(_OUT_EN)
    print(f"[OK] {_OUT}")
    print(f"[OK] {_OUT_EN}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
