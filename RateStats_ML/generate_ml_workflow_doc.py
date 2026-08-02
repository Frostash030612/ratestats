#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《RateStats_ML 选链流程说明》Word 文档（含单家银行流程图）。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError:
    raise SystemExit("请先安装: pip install python-docx matplotlib")

_ML = Path(__file__).resolve().parent
_ROOT = _ML.parent
_DOCS = _ML / "docs"
_OUT = _DOCS / "RateStats_ML选链流程说明.docx"
_FLOW_PNG = _DOCS / "单家银行ML选链流程图.png"


def _h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(10.5)


def _bullet(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _mono(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            t.rows[i + 1].cells[j].text = str(cell)


def _draw_single_bank_flowchart(path: Path) -> None:
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    fig, ax = plt.subplots(figsize=(8, 11))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 14)
    ax.axis("off")

    def box(x, y, w, h, text, color="#E8F4FC", edge="#2B6CB0"):
        rect = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.15",
            linewidth=1.2,
            edgecolor=edge,
            facecolor=color,
        )
        ax.add_patch(rect)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, wrap=True)

    def arrow(x1, y1, x2, y2):
        ax.annotate(
            "",
            xy=(x2, y2),
            xytext=(x1, y1),
            arrowprops=dict(arrowstyle="->", color="#4A5568", lw=1.2),
        )

    ax.text(5, 13.4, "单家银行 ML 选链整体流程（以 OCBC 为例）", ha="center", fontsize=12, weight="bold")

    box(1.5, 12.0, 7, 0.9, "手动维护 url_params.xlsx\nOCBC 新元促销 / 新币挂牌 / 外币挂牌 等 dest → 黄金 URL", "#FFF9E6", "#B7791F")
    box(1.5, 10.5, 7, 0.9, "离线训练 train.py\n黄金=正样本；错链+发现候选=负样本 → url_ranker.joblib", "#EDF2F7", "#4A5568")
    box(1.5, 9.0, 7, 0.9, "评估 evaluate_picker.py + 调参 tune_picker.py\n→ picker_tuning.json（融合权重）", "#EDF2F7", "#4A5568")

    arrow(5, 12.0, 5, 11.4)
    arrow(5, 10.5, 5, 9.9)
    arrow(5, 9.0, 5, 8.5)

    box(0.8, 7.0, 8.4, 1.2, "【线上·按 dest 循环】例如 ocbc_board_url\nQuery: ocbc sgd fixed deposit interest rates", "#E6FFFA", "#2C7A7B")

    box(1.2, 5.5, 3.6, 1.0, "Vertex 搜索\ntop 候选 URL 列表", "#EBF8FF", "#2B6CB0")
    box(5.2, 5.5, 3.6, 1.0, "pick_best_url_ml\n规则分 + ML 概率 + 调参\n硬拒绝 / fallback", "#EBF8FF", "#2B6CB0")

    arrow(5, 7.0, 5, 6.6)
    arrow(3.0, 5.5, 3.0, 5.0)
    arrow(7.0, 5.5, 7.0, 5.0)
    arrow(4.8, 6.0, 3.2, 6.0)
    arrow(5.2, 6.0, 6.8, 6.0)

    box(1.5, 3.8, 7, 0.9, "写入 runs/日期/url_params_ai_ml.xlsx\n+ ai_search_discovered_ml_*.xlsx", "#F0FFF4", "#276749")
    box(1.5, 2.3, 7, 1.0, "bank_all_promo_rates 抓取 OCBC 页面\n解析表格 → Market 中 OCBC 行", "#F0FFF4", "#276749")
    box(1.5, 0.8, 7, 1.0, "compare_ml_market.py\n与同日手动 Market 对比\n（促销 / 挂牌 / 挂牌_USD 等 sheet）", "#FFF5F5", "#C53030")

    arrow(3.0, 5.5, 5, 4.7)
    arrow(7.0, 5.5, 5, 4.7)
    arrow(5, 3.8, 5, 3.3)
    arrow(5, 2.3, 5, 1.8)

    ax.text(
        5,
        0.2,
        "说明：每个 dest 独立搜链；同一银行可有多个 dest（促销、挂牌、外币）",
        ha="center",
        fontsize=8,
        color="#718096",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build() -> Document:
    doc = Document()
    title = doc.add_heading("RateStats_ML 选链流程说明", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p(doc, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    _p(doc, "项目路径：RateStats/RateStats_ML（独立实验，不覆盖 Portable 规则版 AI 流程）")

    _h(doc, "一、ML 方案在干什么？", 1)
    _p(
        doc,
        "目标：在 Vertex 搜索返回的多条候选 URL 中，选出最接近手动维护「黄金链接」的一条，"
        "再用于抓取 MarketRateData。不替代网页解析；不修改手动 url_params.xlsx；"
        "产出独立的 url_params_ai_ml.xlsx，输出目录为 runs/YYYYMMDD/。",
    )

    _h(doc, "二、与规则版 Vertex 的对比", 1)
    _table(
        doc,
        ["环节", "规则版（Portable）", "ML 版（RateStats_ML）"],
        [
            ["搜索", "Vertex API", "相同 Vertex API"],
            ["选链", "pick_best_url", "pick_best_url_ml（规则+模型+调参）"],
            ["URL 配置", "assets/url_params_ai.xlsx", "runs/<日期>/url_params_ai_ml.xlsx"],
            ["Market", "MarketRateData_*_AISearch.xlsx", "MarketRateData_*_ML.xlsx"],
            ["训练", "无", "本地 sklearn，无 API 费"],
        ],
    )

    _h(doc, "三、推荐操作顺序（四步）", 1)
    _table(
        doc,
        ["步骤", "批处理", "作用"],
        [
            ["1", "01_训练模型.bat", "train.py --use-discovery：黄金标签 + runs 发现报告 + 错链负样本"],
            ["2", "02_评估选链准确率.bat", "evaluate_picker.py：总体与 by_dest 命中率"],
            ["3", "04_调参融合权重.bat", "tune_picker.py → models/picker_tuning.json"],
            ["4", "03_Vertex_ML发现并抓取.bat", "发现 + 抓取 → runs/日期/"],
        ],
    )
    _mono(
        doc,
        "cd RateStats_ML\npython train.py --use-discovery\npython evaluate_picker.py\n"
        "python tune_picker.py\npython run_vertex_ml_discovery.py --run-tag 20260603 --fetch",
    )

    _h(doc, "四、阶段 A：离线（训练 / 评估 / 调参）", 1)

    _h(doc, "4.1 标签与训练数据", 2)
    _bullet(
        doc,
        [
            "黄金标准：RateStats_Portable/assets/url_params.xlsx 每个 dest 的手动 URL。",
            "正样本：该 dest 的黄金 URL（label=1）。",
            "负样本：其它 dest 的黄金 URL、路径扰动、hard_negatives 典型错链、发现报告候选（与黄金比对为 0）。",
            "发现报告目录：assets/、runs/**、AI_Compare/output/（ai_search_discovered*.xlsx）。",
            "可编辑 data/hard_negatives.csv 追加错链后重新训练。",
        ],
    )

    _h(doc, "4.2 特征与模型", 2)
    _p(doc, "features.py 抽取约 27 维特征：规则总分及 path/intent 子项、与黄金链相似度、dest 类型、错页模式（premier、business、API 等）。")
    _p(doc, "train.py 使用 HistGradientBoostingClassifier，样本权重 balanced；产出 url_ranker.joblib 与 url_ranker_meta.json。")

    _h(doc, "4.3 评估与调参", 2)
    _p(doc, "evaluate_picker.py 在历史发现报告上回放规则 vs ML，输出 picker_eval_*.xlsx（summary / by_dest / baseline_miss_dest）。")
    _p(doc, "tune_picker.py 网格搜索 blend_rule_weight、min_ml_proba，写入 picker_tuning.json；推理时 picker.py 自动读取。")
    _p(doc, "融合公式：综合分 = blend × 规则分(归一化) + (1-blend) × ML 概率；规则分差大时跳过 ML（ml_skip_clear_rule）。")

    _h(doc, "五、阶段 B：线上（发现 + 抓取）", 1)
    _bullet(
        doc,
        [
            "入口：run_vertex_ml_discovery.py（或 03 批处理）。",
            "对每个 dest：构造 query → Vertex 搜索 → pick_best_url_ml 选链。",
            "过滤：银行域名、url_hard_reject；低置信回退规则或 fallback。",
            "写出：url_params_ai_ml.xlsx、ai_search_discovered_ml_*.xlsx。",
            "--fetch：调用 bank_all_promo_rates 生成 MarketRateData_*_ML.xlsx。",
            "对比：compare_ml_market.py 与同日手动 Market 比较。",
        ],
    )

    _h(doc, "六、目录与产物", 1)
    _mono(
        doc,
        r"""RateStats/
├── runs/YYYYMMDD/
│   ├── MarketRateData_*_manual.xlsx
│   ├── MarketRateData_*_ML.xlsx
│   ├── url_params_ai_ml.xlsx
│   └── market_data_compare_*_ML.xlsx
├── RateStats_Portable/assets/url_params.xlsx
└── RateStats_ML/
    ├── models/url_ranker.joblib
    ├── models/picker_tuning.json
    └── output/picker_eval_*.xlsx""",
    )

    _h(doc, "七、能力边界", 1)
    _bullet(
        doc,
        [
            "候选列表中没有黄金 URL 时，规则与 ML 均只能 fallback，无法猜对。",
            "Market 利率一致还依赖同日抓取、页面改版、解析逻辑；选链只是第一步。",
            "训练/推理无 Vertex 费用；仅 03 发现+抓取消耗搜索 API。",
            "难点 dest（BOC、RHB PDF、DBS 外币 API 等）需错链样本与规则并行加强。",
        ],
    )

    _h(doc, "八、单家银行整体流程图", 1)
    _p(doc, "下图以 OCBC 为例：同一银行在 url_params 中对应多个 dest（促销、新币挂牌、外币挂牌），每个 dest 独立走「搜索 → ML 选链 → 抓取 → 对比」。")

    if _FLOW_PNG.is_file():
        doc.add_picture(str(_FLOW_PNG), width=Inches(5.8))
    else:
        _p(doc, "（流程图 PNG 未生成，请运行 generate_ml_workflow_doc.py）")

    _h(doc, "九、一句话总结", 1)
    _p(
        doc,
        "用手动 URL 当标准答案 → 训练「候选是否像黄金链」的模型并调融合参数 → "
        "Vertex 搜链后用 ML+规则选出 URL → 抓取 Market → 与手动表对比验证。",
    )

    return doc


def main() -> int:
    _DOCS.mkdir(parents=True, exist_ok=True)
    _draw_single_bank_flowchart(_FLOW_PNG)
    doc = build()
    doc.save(_OUT)
    print(f"[OK] {_OUT}")
    print(f"[OK] {_FLOW_PNG}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
