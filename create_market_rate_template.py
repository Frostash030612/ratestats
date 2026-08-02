#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Create a fixed Excel template for MarketRateData output.

This template is intended to be edited only when the *layout* changes
(new tenor columns, changed sheet names/row offsets, etc).

Runtime requirement:
- bank_all_promo_rates.write_rates_excel will copy this template to a temp
  working file and only write *values* into the data area.

Currency sheets:
- The template includes a base sheet `挂牌_USD` with the full FX board layout.
- When new currencies appear in data, the script auto-copies `挂牌_USD` to
  create `挂牌_{币种}` sheets.
- Only when the FX board layout itself changes (column headers/tenors),
  you should re-run this generator to update the template.
"""

from __future__ import annotations

import os
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def _header_font() -> Font:
    return Font(bold=True)


def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center")


def _safe_sheet_title(title: str) -> str:
    # Excel constraints: max 31 chars, cannot contain: \ / ? * : [ ]
    s0 = "".join("_" if c in "\\/*?:[]" else c for c in (title or "").strip())
    s0 = s0[:31] if len(s0) > 31 else s0
    return s0 or "Sheet"


def _setup_table_header(ws, headers: List[str], row_idx_1based: int = 1) -> None:
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx_1based, column=i, value=h)
        cell.font = _header_font()
        cell.alignment = _center()


def _setup_fx_board_sheet(
    ws,
    title: str,
    headers: List[str],
) -> None:
    # Layout:
    # - Row 1: merged big title
    # - Row 2: column headers
    ncols = len(headers)
    end_col = get_column_letter(ncols)

    ws.merge_cells(f"A1:{end_col}1")
    a1 = ws["A1"]
    a1.value = title
    a1.font = Font(size=16, bold=True)
    a1.alignment = _center()
    ws.row_dimensions[1].height = 30

    _setup_table_header(ws, headers, row_idx_1based=2)


def main() -> None:
    base_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(base_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)
    out_path = os.path.join(assets_dir, "MarketRateData_template.xlsx")

    # Column headers must match bank_all_promo_rates.write_rates_excel.
    meta_headers = ["项目", "内容"]

    sgd_col_order = [
        "数据来源",
        "产品或客群",
        "起存金额_页面",
        "上限金额_页面",
        "资金区间_页面",
        "1M_pct",
        "3M_pct",
        "5M_pct",
        "6M_pct",
        "9M_pct",
        "12M_pct",
        "页面原文摘录",
    ]

    sgd_board_col_order = [
        "数据来源",
        "产品或档位",
        "起存金额_页面",
        "上限金额_页面",
        "资金区间_页面",
        "7D_pct",
        "14D_pct",
        "1M_pct",
        "2M_pct",
        "3M_pct",
        "4M_pct",
        "5M_pct",
        "6M_pct",
        "7M_pct",
        "8M_pct",
        "9M_pct",
        "10M_pct",
        "11M_pct",
        "12M_pct",
        "18M_pct",
        "24M_pct",
        "36M_pct",
        "页面原文摘录",
    ]

    fx_time_deposit_headers = [
        "币种",
        "1M_pct",
        "3M_pct",
        "6M_pct",
        "9M_pct",
        "12M_pct",
        "18M_pct",
        "24M_pct",
        "起存金额_页面",
        "上限金额_页面",
        "数据来源",
        "页面_1M原文",
    ]

    other_headers = ["类别", "指标", "数值"]
    fx_notes_headers = ["说明"]

    fx_board_col_order = [
        "数据来源",
        "币种",
        "产品或档位",
        "起存金额_页面",
        "上限金额_页面",
        "资金区间_页面",
        "7D_pct",
        "14D_pct",
        "1M_pct",
        "2M_pct",
        "3M_pct",
        "6M_pct",
        "9M_pct",
        "12M_pct",
        "18M_pct",
        "24M_pct",
        "36M_pct",
        "48M_pct",
        "60M_pct",
        "页面原文摘录",
    ]

    # Sheet titles
    meta_ws = "元数据"
    sgd_promo_ws = "新元定存促销"
    sgd_board_ws = "新元挂牌利率"
    fx_promo_ws = "外币定存促销"
    fx_notes_ws = "外币定存备注"
    other_ws = "其它利率"

    fx_base_ws = "挂牌_USD"  # base layout to copy when new currencies appear

    # 预置的外币挂牌 sheet：
    # - `外币挂牌利率` 汇总占位页不再预置（数据按币种分散到 `挂牌_*`）。
    # - SGD 不再预置外币挂牌分 sheet（SGD 走「新元挂牌利率」）。
    prebuilt_currencies = [
        "USD",
        "AUD",
        "EUR",
        "GBP",
        "HKD",
        "JPY",
        "CHF",
        "NZD",
        "CAD",
        "CNH",
    ]

    wb = Workbook()
    # Remove the default sheet (keep it only if we end up not creating a replacement).
    default_ws = wb.active
    default_ws.title = "TEMP_REMOVE"

    def add_sheet(title: str):
        title = _safe_sheet_title(title)
        ws = wb.create_sheet(title=title)
        return ws

    # Main sheets
    ws_meta = add_sheet(meta_ws)
    _setup_table_header(ws_meta, meta_headers, row_idx_1based=1)

    ws_sgd = add_sheet(sgd_promo_ws)
    _setup_table_header(ws_sgd, sgd_col_order, row_idx_1based=1)

    ws_sgd_board = add_sheet(sgd_board_ws)
    _setup_table_header(ws_sgd_board, sgd_board_col_order, row_idx_1based=1)

    ws_fx = add_sheet(fx_promo_ws)
    _setup_table_header(ws_fx, fx_time_deposit_headers, row_idx_1based=1)

    ws_fx_notes = add_sheet(fx_notes_ws)
    _setup_table_header(ws_fx_notes, fx_notes_headers, row_idx_1based=1)

    ws_other = add_sheet(other_ws)
    _setup_table_header(ws_other, other_headers, row_idx_1based=1)

    # Base currency layout sheet
    ws_fx_base = add_sheet(fx_base_ws)
    _setup_fx_board_sheet(ws_fx_base, "USD 外币挂牌利率", fx_board_col_order)

    # Prebuilt currencies
    for cur in prebuilt_currencies:
        name = _safe_sheet_title(f"挂牌_{cur}")
        if name == fx_base_ws:
            continue
        ws = add_sheet(name)
        _setup_fx_board_sheet(ws, f"{cur} 外币挂牌利率", fx_board_col_order)

    # Remove temp sheet
    wb.remove(wb["TEMP_REMOVE"])

    wb.save(out_path)
    print(f"[OK] Template created: {out_path}")


if __name__ == "__main__":
    main()

