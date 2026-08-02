#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate spoken script for IA_Presentation_RateStats_Xu_Wenzhe.pptx (~10 min).

Narrative: Manual → Python application → AI Search plugged into that flow.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "docs" / "IA_Reports"
OUT.mkdir(parents=True, exist_ok=True)


def _h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str, *, bold: bool = False) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(11)
    run.bold = bold


def _note(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _bullet(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(item, style="List Bullet")
        for run in para.runs:
            run.font.size = Pt(11)


def build() -> Path:
    doc = Document()

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("RateStats — Presentation Script (~10 minutes)")
    r.bold = True
    r.font.size = Pt(18)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(
        "Xu Wenzhe | Bank of China Singapore Branch | FinTech\n"
        "Matches: IA_Presentation_RateStats_Xu_Wenzhe.pptx\n"
        "Story: Manual → Python app → AI Search in the same pipeline"
    )
    sr.font.size = Pt(11)

    _note(
        doc,
        "Timing: ~8–9 minutes speaking + up to 2 minutes demo. "
        "Key message: AI was added later into a Python product you built — "
        "not ‘we started with AI’.",
    )

    _h(doc, "Slide 1 — Title (≈ 20 sec)", 2)
    _p(
        doc,
        "Good afternoon everyone. My name is Xu Wenzhe. Today I will present RateStats — "
        "a market interest rate automation project from my industrial attachment at "
        "Bank of China Singapore Branch, in the Financial Technology department.",
    )
    _p(
        doc,
        "The short story of this project is: it started from fully manual work, I first "
        "built a practical Python application, and later we added AI Search into that "
        "same end-to-end flow.",
    )

    _h(doc, "Slide 2 — Agenda (≈ 20 sec)", 2)
    _p(
        doc,
        "I will cover the project background and how it evolved, my role, the development "
        "phases, project status, business value with a before-and-after comparison, and "
        "a short demo if time allows.",
    )

    _h(doc, "Slide 3 — Project Background (≈ 80 sec)", 2)
    _p(
        doc,
        "Our team needs deposit and promotional rates from major Singapore banks — "
        "about seventeen banks, roughly forty-seven rate nodes. For each bank we may "
        "need SGD board rates, SGD promotional rates, and the same for foreign currency.",
    )
    _p(
        doc,
        "At the beginning, this was fully manual. Colleagues opened bank websites, "
        "copied rates, and maintained Excel by hand. That was slow and easy to miss updates.",
    )
    _p(
        doc,
        "My first job on RateStats was therefore to write Python scripts and turn that "
        "manual work into a real application: fetch pages, parse HTML, PDF, or APIs, "
        "and generate MarketRateData and Rainbow Table Excel. At that stage, URLs were "
        "still provided manually in url_params.xlsx. When a bank redesigned its site, "
        "someone still had to find and update the link.",
    )
    _p(
        doc,
        "So the next stage was to add AI Search into this existing pipeline — discover "
        "candidate URLs automatically, validate them, then reuse the Python extractors "
        "I had already built. In a bank we keep this controlled: AI search and scoring, "
        "not unrestricted cloud LLM agents.",
    )

    _h(doc, "Slide 4 — My Project Role (≈ 55 sec)", 2)
    _p(
        doc,
        "My role started as a Python developer building a usable product, not only an "
        "AI experiment. I implemented the bank extractors and the Excel generation "
        "that business users actually run.",
    )
    _p(
        doc,
        "After that foundation worked, I extended the same pipeline with Vertex AI "
        "Search, compared Serper, Brave, and Tavily, and added hybrid URL selection — "
        "domain rules plus an ML ranker — because AI candidates can look right but be wrong. "
        "I also packaged batch scheduling and URL health checks.",
    )
    _p(
        doc,
        "I also supported RPA bots for other operations, but this presentation focuses "
        "on RateStats.",
    )

    _h(doc, "Slide 5 — Requirements & Design (≈ 60 sec)", 2)
    _p(
        doc,
        "Requirements came in two phases. Phase A: replace hand-copying of rates with "
        "reliable Python fetch, parse, and Excel output. That meant bank-specific "
        "parsers for HTML, PDF tables, and JSON APIs.",
    )
    _p(
        doc,
        "Phase B: stop feeding every URL by hand. Use AI Search to find the latest "
        "pages, and reject wrong links — wrong bank domain, Premier versus board page, "
        "PDF versus HTML. The final design plugs Discovery and Scoring in front of "
        "the Extract and Output engine that already existed.",
    )
    _p(
        doc,
        "For example, ICBC alone needs four page types — SGD board, SGD promo, FCY "
        "board, FCY promo — and we repeat that pattern across the bank list.",
    )

    _h(doc, "Slide 6 — Implementation / Architecture (≈ 70 sec)", 2)
    _p(
        doc,
        "This slide shows the end-to-end architecture. Steps three and four — Extract "
        "and Output — were built first as the Python application. Steps one and two — "
        "Discover and Score — were added later.",
    )
    _p(
        doc,
        "Today the flow is: AI Search returns candidates; rules and ML choose a valid "
        "URL; Python parsers extract rates; Excel is written for business users.",
    )
    _p(
        doc,
        "A real scenario after AI was added: Search sometimes returned an HLF page "
        "for an ICBC destination. Both talk about fixed deposits, so it looked "
        "plausible — but it was the wrong bank. We fixed that with a per-destination "
        "host whitelist before ML scoring.",
    )

    _h(doc, "Slide 7 — Testing & Deployment (≈ 45 sec)", 2)
    _p(
        doc,
        "We still keep manual golden URLs as the ground truth to test AI picks. We "
        "evaluate the ML picker, benchmark four search providers, and run health "
        "checks for HTTP errors and redirects. Deployment is batch scripts plus "
        "Task Scheduler into dated runs folders, delivering MarketRateData and the "
        "Rainbow Table.",
    )

    _h(doc, "Slide 8 — Project Status (≈ 40 sec)", 2)
    _p(
        doc,
        "The Python extraction and Excel pipeline — the foundation — are completed "
        "and in production. AI discovery, rule scoring, and multi-provider "
        "benchmarking are also completed. The ML hybrid ranker is integrated and "
        "still being tuned, so I mark that piece as in progress for production.",
    )

    _h(doc, "Slide 9 — Business Value Before / After (≈ 70 sec)", 2)
    _p(
        doc,
        "Before: fully manual collection — open sites, copy rates, heavy manpower "
        "across about forty-seven nodes, easy to miss updates.",
    )
    _p(
        doc,
        "After: a Python application generates the Excel, and AI Search plus "
        "validation maintains the URLs in the same flow. Runs can be scheduled; "
        "missing rows are fewer. On evaluation, URL hit rate is about ninety-five "
        "percent, and the benchmark supported choosing Vertex. The tangible saving "
        "is manpower and more reliable market comparison reports.",
    )
    _p(
        doc,
        "Please note the middle step as well: even before AI, the Python scripts "
        "already removed most of the hand-copying of rates. AI mainly removed the "
        "remaining manual URL hunting.",
    )

    _h(doc, "Slide 10 — Value Added by the Intern (≈ 45 sec)", 2)
    _p(
        doc,
        "What I added was delivering a usable Python product first, growing scripts "
        "into a schedulable application, then plugging AI Search into that real "
        "pipeline — not a separate demo. When pure AI candidates failed, I designed "
        "hybrid rules plus ML, and the vendor benchmark turned a preference into "
        "an evidence-based decision for the team.",
    )

    _h(doc, "Slide 11 — Demo (≤ 2 min)", 2)
    _note(doc, "[If no live demo, skip this slide and jump to closing.]")
    _p(
        doc,
        "I will show a short demo under two minutes: run the batch that discovers "
        "URLs and fetches rates through the Python extractors, briefly show a wrong "
        "candidate being rejected, then open the generated Excel.",
    )

    _h(doc, "Slide 12 — Conclusion & Q&A (≈ 30 sec)", 2)
    _p(
        doc,
        "To conclude: RateStats went from fully manual collection, to a Python "
        "automation application, to AI Search inside that same flow. AI is an "
        "extension of a working product, not the starting point. Hybrid rules and "
        "ML keep URL accuracy high under bank security constraints. The business "
        "impact is less manpower and more reliable market-rate Excel. Thank you. "
        "I am happy to take your questions.",
    )

    doc.add_page_break()
    _h(doc, "Full Continuous Script (read-through version)", 1)
    _note(doc, "Practice as one speech without looking at slide titles.")

    full = """Good afternoon everyone. My name is Xu Wenzhe. Today I will present RateStats — a market interest rate automation project from my industrial attachment at Bank of China Singapore Branch, in the Financial Technology department. The short story is: it started from fully manual work; I first built a practical Python application; later we added AI Search into that same end-to-end flow.

I will cover background, my role, development phases, status, business value, and a short demo if time allows.

We need deposit and promotional rates from about seventeen Singapore banks — roughly forty-seven rate nodes. At the beginning this was fully manual: open websites, copy rates, maintain Excel by hand. My first job was to write Python scripts and turn that into a real application — fetch, parse HTML, PDF or APIs, and generate MarketRateData and Rainbow Table Excel. URLs were still fed manually in url_params.xlsx, so when a bank redesigned its site someone still had to update the link. The next stage was to add AI Search into this existing pipeline: discover candidates, validate them, then reuse the Python extractors. In a bank we keep this controlled — search and scoring, not unrestricted cloud LLM agents.

My role therefore started as a Python developer building a usable product. After the foundation worked, I extended it with Vertex AI Search, compared other providers, and added hybrid URL selection with domain rules and an ML ranker. I also packaged scheduling and health checks. I did RPA work as well, but today I focus on RateStats.

Requirements came in two phases. Phase A: automate fetch, parse, and Excel. Phase B: stop hand-feeding every URL — use AI Search and reject wrong links. The final design plugs Discovery and Scoring in front of the Extract and Output engine that already existed.

On the architecture slide, steps three and four were built first; steps one and two were added later. Today AI proposes URLs, rules and ML validate, Python extracts, Excel is delivered. Example: AI once returned HLF for ICBC; we blocked that with a host whitelist.

We test against golden URLs, evaluate the picker, benchmark four providers, and run health checks. Deployment uses Task Scheduler and dated run folders.

Python extraction and Excel are in production. AI discovery and rule scoring are completed. The ML ranker is integrated and still being tuned.

Before: fully manual, heavy manpower. After: Python app plus AI Search in the same flow — scheduled runs, about ninety-five percent URL hit rate, Vertex chosen by benchmark. Importantly, even before AI, Python already removed most hand-copying of rates; AI mainly removed remaining manual URL hunting.

What I added was a usable Python product first, then AI plugged into that real pipeline, plus hybrid rules and ML when pure AI failed, and evidence-based provider comparison.

If time allows I will demo the batch and the Excel briefly.

To conclude: RateStats went from manual, to Python automation, to AI Search in the same flow. AI is an extension of a working product, not the starting point. Thank you. I am happy to take your questions."""

    for para in full.strip().split("\n\n"):
        _p(doc, para)

    doc.add_page_break()
    _h(doc, "中文提词卡（排练用）", 1)
    _bullet(
        doc,
        [
            "核心叙事：全人工 → 我先写 Python 做成可用应用 → 再把 AI 搜索接进同一条流水线",
            "背景：约 17 家银行 / 47 节点；最初人工开网页抄利率",
            "阶段一：Python 抓取 HTML/PDF/API + 出 MarketRateData/彩虹表；URL 仍手动配",
            "阶段二：AI Search 找链 + 规则/ML 校验 → 复用已有解析器",
            "角色：先做产品再加 AI，不是一上来就做大模型",
            "架构：3–4 步先做，1–2 步后加；ICBC→HLF 用域名白名单",
            "价值：Before 全人工；After Python+AI；中间态已去掉大部分手抄利率",
            "收尾：AI 是扩展，不是起点；谢谢提问",
        ],
    )

    _h(doc, "时间不够时的删减顺序", 2)
    _bullet(
        doc,
        [
            "先砍 Demo",
            "再压缩 Testing（一句带过）",
            "务必保留：Background 三阶段演进、Architecture「先 Python 后 AI」、Before/After",
        ],
    )

    path = OUT / "IA_Presentation_RateStats_Script_Xu_Wenzhe.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build())
