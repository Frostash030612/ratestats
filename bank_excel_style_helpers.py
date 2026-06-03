from __future__ import annotations

"""
Excel 样式辅助模块（openpyxl）。

这个模块专门承载 `bank_all_promo_rates.py` 里 `write_rates_excel()` 使用到的
“样式 helper”逻辑，避免把大量 openpyxl 样式细节塞在主函数内部（嵌套 def）。

抽取后主脚本会直接调用这些函数，行为应与原先嵌套版保持一致。
"""

import math
from copy import copy
import sys
from typing import Any, List, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# Light pastel palette by bank/source (same source keeps same color across sheets).
BANK_FILL_MAP: dict[str, str] = {
    "CITIBANK": "FFFDE7",
    "CIMB": "E8F5E9",
    "DBS": "FCE4EC",
    "HL BANK": "E3F2FD",
    "HLF": "E1F5FE",
    "HSBC": "FBE9E7",
    "ICBC": "EDE7F6",
    "MAYBANK": "FFF9C4",
    "OCBC": "F3E5F5",
    "RHB": "F1F8E9",
    "SINGFINANCE": "E0F2F1",
    "SCB": "E8EAF6",
    "SBI": "E0F7FA",
    "UOB": "FFF3E0",
    "BOC": "F9FBE7",
    "BEA": "F5F5F5",
}


def pick_bank_fill(source_value: Any) -> Optional[PatternFill]:
    """
    根据单元格“数据来源”文本选取对应的浅色填充（PatternFill）。

    source_value:
        通常是工作表里某一行的“数据来源”字段值（字符串，可能包含银行名字）。

    返回：
        匹配到的 PatternFill；匹配不到则返回 None。
    """
    try:
        src = str(source_value or "").upper()
        for key, color in BANK_FILL_MAP.items():
            if key in src:
                return PatternFill(fill_type="solid", start_color=color, end_color=color)
        return None
    except Exception as e:
        # 样式失败不应中断数据输出
        print(f"Excel 样式：pick_bank_fill 失败：{e}", file=sys.stderr)
        return None


def apply_bank_colors(
    ws: Any,
    source_col_name: str = "数据来源",
    header_row: int = 1,
) -> None:
    """
    给工作表中“数据来源”匹配到的行设置浅色填充。

    逻辑与原先 write_rates_excel() 内嵌版保持一致：
    1) 找到 header_row 这一行里名为 source_col_name 的列索引；
    2) 遍历 header_row+1 .. max_row，对每一行读取该列值；
    3) 用 pick_bank_fill() 得到填充色，若不为 None 则整行每个单元格填充。
    """
    try:
        if ws.max_row < header_row:
            return

        hdr_cells = list(ws[header_row])
        src_idx = None
        for i, c in enumerate(hdr_cells, start=1):
            if str(c.value or "").strip() == source_col_name:
                src_idx = i
                break
        if src_idx is None:
            return

        for r in range(header_row + 1, ws.max_row + 1):
            src_val = ws.cell(row=r, column=src_idx).value
            fill = pick_bank_fill(src_val)
            if not fill:
                continue
            for cc in range(1, ws.max_column + 1):
                ws.cell(row=r, column=cc).fill = fill
    except Exception as e:
        print(f"Excel 样式：apply_bank_colors 失败：{e}", file=sys.stderr)
        return


def center_align_rate_columns(ws: Any, header_row: int = 1) -> None:
    """
    把所有“利率列”（表头以 `_pct` 结尾）下方的单元格居中对齐。

    适用场景：
        - 元数据/促销/挂牌表中利率列一般命名为 `*_pct`。
    """
    try:
        if ws.max_row < header_row:
            return

        hdr_cells = list(ws[header_row])
        rate_col_idxs: List[int] = []
        for i, c in enumerate(hdr_cells, start=1):
            h = str(c.value or "").strip()
            if h.endswith("_pct"):
                rate_col_idxs.append(i)
        if not rate_col_idxs:
            return

        center = Alignment(horizontal="center", vertical="center")
        for ci in rate_col_idxs:
            for r in range(header_row + 1, ws.max_row + 1):
                ws.cell(row=r, column=ci).alignment = center
    except Exception as e:
        print(f"Excel 样式：center_align_rate_columns 失败：{e}", file=sys.stderr)
        return


def style_sgd_board_sheet(ws: Any) -> None:
    """
    专门的“新元挂牌利率”sheet 样式。

    样式包括：
    1) header 行（第 1 行）设置深蓝底色 + 白色加粗字体；
    2) `*_pct` 列的数据居中；
    3) 表格外边框粗线（thick），内边框细线（thin）；
    4) 每行里 `*_pct` 的最大值对应的单元格字体加粗（若同最大值存在多列也会同时加粗）。
    """
    try:
        if ws.max_row < 1 or ws.max_column < 1:
            return

        max_r = ws.max_row
        max_c = ws.max_column

        thin = Side(style="thin", color="000000")
        thick = Side(style="thick", color="000000")
        header_fill = PatternFill(fill_type="solid", start_color="1F4E79", end_color="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Header row styling (row 1)
        for c in range(1, max_c + 1):
            cell = ws.cell(1, c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

    # Locate rate columns: header names ending with "_pct"
        pct_cols: List[int] = []
        for c in range(1, max_c + 1):
            h = str(ws.cell(1, c).value or "").strip()
            if h.endswith("_pct"):
                pct_cols.append(c)

    # Align entire grid center (row 2..max_r)
        for r in range(2, max_r + 1):
            for c in range(1, max_c + 1):
                ws.cell(r, c).alignment = center

    # Apply borders with thick outer frame
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                left = thick if c == 1 else thin
                right = thick if c == max_c else thin
                top = thick if r == 1 else thin
                bottom = thick if r == max_r else thin
                ws.cell(r, c).border = Border(left=left, right=right, top=top, bottom=bottom)

    # Bold best value per row (max among *_pct columns)
        for r in range(2, max_r + 1):
            nums: List[tuple[int, float]] = []
            for c in pct_cols:
                v = ws.cell(r, c).value
                if v is None:
                    continue
                if isinstance(v, str) and str(v).strip().upper() in ("N/A", "NA", ""):
                    continue
                try:
                    fv = float(v)
                except (TypeError, ValueError):
                    continue
                if math.isnan(fv):
                    continue
                nums.append((c, fv))

        # Less than 2 numeric points: keep font unchanged (matches original behavior).
            if len(nums) < 2:
                continue

            best = max(f for _, f in nums)
            for c, fv in nums:
                if abs(fv - best) > 1e-12:
                    continue
                cell = ws.cell(r, c)
                nf = copy(cell.font)
                nf.bold = True
                cell.font = nf
    except Exception as e:
        print(f"Excel 样式：style_sgd_board_sheet 失败：{e}", file=sys.stderr)
        return


# Safety net: 自动给本模块中所有 def 加异常兜底
from bank_safety import auto_wrap_module_functions as _auto_wrap_module_functions

_auto_wrap_module_functions(globals(), module_name=__name__)

