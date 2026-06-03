from __future__ import annotations

"""
MarketRateData 模板写入器。

把 `bank_all_promo_rates.py` 里 `write_rates_excel()` 的“模板复制 + 按币种拆 sheet + overlay 写入 + 应用样式”
这一整段逻辑抽离出来，避免主脚本 write_rates_excel() 过长。
"""

import os
import shutil
from datetime import datetime
from typing import Any, List, Optional, Tuple

import pandas as pd

from bank_excel_style_helpers import (
    apply_bank_colors as _apply_bank_colors,
    center_align_rate_columns as _center_align_rate_columns,
    style_sgd_board_sheet as _style_sgd_board_sheet,
)

from bank_safety import safe_call


def _fx_board_currency_label(cur: Any) -> str:
    """
    将 payload 里的“币种值”归一为用于 sheet 标题的短标签。

    主要作用：
    1) 避免 None / 空字符串导致 sheet 名称不稳定；
    2) 将 NaN 视为“未填币种”，统一落到一个可预期的标题后缀。
    """
    if cur is None:
        return "（未填币种）"
    if isinstance(cur, float):
        # NaN -> 未填币种
        try:
            if pd.isna(cur):
                return "（未填币种）"
        except Exception:
            pass
    t = str(cur).strip()
    return t if t else "（未填币种）"


def _is_sgd_fx_board_currency(cur: Any) -> bool:
    """
    判断该行币种是否应视为 SGD。

    SGD 挂牌数据应落在「新元挂牌利率」sheet，不单独生成「挂牌_SGD / SGD 外币挂牌利率」。
    """
    return _fx_board_currency_label(cur).strip().upper() == "SGD"


def _remove_sheet_if_exists(wb: Any, name: str) -> None:
    """从工作簿中删除指定 sheet（存在则删，不存在则忽略）。"""
    if name in wb.sheetnames:
        wb.remove(wb[name])


@safe_call(
    context="写入 MarketRateData Excel（模板覆写）",
    default_value=None,
)
def write_market_rate_excel_with_template(
    *,
    df_meta: pd.DataFrame,
    df_sgd: pd.DataFrame,
    df_sgd_board: pd.DataFrame,
    df_fx: pd.DataFrame,
    df_fx_notes: pd.DataFrame,
    df_other: pd.DataFrame,
    df_fx_board: pd.DataFrame,
    path: str,
    script_dir: str,
) -> None:
    """
    使用 `assets/MarketRateData_template.xlsx` 作为固定模板，将传入的 DataFrame
    覆盖/写入到正确的 sheet 与区域（不覆盖模板表头）。

    模板/写入约定（实现中依赖这些名字/行号）：
    - sheet 名称是固定的：`元数据`、`新元定存促销`、`新元挂牌利率`、`外币定存促销`、`外币定存备注`、`其它利率`。
    - FX board（外币挂牌）会根据 `df_fx_board['币种']` 创建/复用 sheet，命名为 `挂牌_{标签}`；
      **不包含 SGD**（SGD 走「新元挂牌利率」）。
    - 模板中的汇总占位 sheet `外币挂牌利率` 不再输出；模板若含 `挂牌_SGD` 也会移除（数据已按币种分散到 `挂牌_*`，SGD 见「新元挂牌利率」）。
    - 写入时使用 `header=False`，并通过固定 `startrow` 只覆盖“数据区”，避免把模板里的表头覆盖掉。

    该函数负责的步骤：
    1) 找到模板文件 -> 复制到 `temp/` -> `openpyxl.load_workbook`
    2) 根据 `df_fx_board['币种']` 创建/命名 FX 多币种 sheet，并设置每张 sheet 的 A1 大标题
    3) 用 `pd.ExcelWriter(engine='openpyxl', if_sheet_exists='overlay')` 只写数据区（header=False）
    4) 对“新元定存促销/外币定存促销/新元挂牌利率”等 sheet 应用配色/居中/边框样式
    5) 复制临时文件到最终 `path`
    """

    # -----------------------------------------------------------------------
    # 0) 找模板 + 复制到 temp 工作区
    # -----------------------------------------------------------------------
    # 注意：这里不直接在原模板文件上写，而是 copy 到临时文件，避免破坏版本化模板。
    template_path_candidates = [
        os.path.join(script_dir, "assets", "MarketRateData_template.xlsx"),
        os.path.join(script_dir, "..", "assets", "MarketRateData_template.xlsx"),
    ]
    template_path = next(
        (p for p in template_path_candidates if os.path.exists(p)), None
    )
    if not template_path:
        raise FileNotFoundError(
            "Cannot find MarketRateData_template.xlsx in assets/. "
            f"Looked in: {template_path_candidates}"
        )

    temp_dir = os.path.join(script_dir, "temp")
    os.makedirs(temp_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    temp_workbook_path = os.path.join(
        temp_dir, f"MarketRateData_work_{ts}.xlsx"
    )
    shutil.copyfile(template_path, temp_workbook_path)

    from openpyxl import load_workbook

    wb = load_workbook(temp_workbook_path)

    # -----------------------------------------------------------------------
    # 1) 处理 FX 多币种 sheet：创建 sheet + 设置标题
    # -----------------------------------------------------------------------
    # - `挂牌_USD` 作为 base sheet（模板里的样式/表头/列宽等都从它复制而来）
    # - 不再保留模板里的汇总 sheet「外币挂牌利率」（无数据占位）：外币挂牌仅按币种分散到 `挂牌_*`
    # - 币种为 SGD 的行不生成外币挂牌 sheet（SGD 挂牌见「新元挂牌利率」）
    fx_board_sheet_names: List[str] = []
    fx_base_sheet = "挂牌_USD"
    fx_empty_sheet = "外币挂牌利率"
    curr_col = "币种"

    fx_groups_sorted: List[Tuple[Any, pd.DataFrame]] = []
    if df_fx_board.empty or curr_col not in df_fx_board.columns:
        _remove_sheet_if_exists(wb, fx_empty_sheet)
    else:
        groups = list(df_fx_board.groupby(curr_col, dropna=False, sort=False))
        groups = [
            (k, v) for k, v in groups if not _is_sgd_fx_board_currency(k)
        ]
        groups.sort(key=lambda kv: _fx_board_currency_label(kv[0]))
        fx_groups_sorted = groups

        if not fx_groups_sorted:
            _remove_sheet_if_exists(wb, fx_empty_sheet)
        else:
            for cur_key, sub_df in fx_groups_sorted:
                label = _fx_board_currency_label(cur_key)
                sn_fb = f"挂牌_{label}"
                # 新币种 sheet：复制模板 base sheet（例如 USD）并重命名
                if sn_fb not in wb.sheetnames:
                    if fx_base_sheet not in wb.sheetnames:
                        raise RuntimeError(
                            f"Template missing base fx board sheet {fx_base_sheet!r} "
                            f"needed to copy for new currency {label!r}."
                        )
                    base_ws = wb[fx_base_sheet]
                    new_ws = wb.copy_worksheet(base_ws)
                    new_ws.title = sn_fb

                wb[sn_fb]["A1"].value = f"{label} 外币挂牌利率"
                fx_board_sheet_names.append(sn_fb)

            _remove_sheet_if_exists(wb, fx_empty_sheet)

    # 模板可能预置「挂牌_SGD」；SGD 挂牌仅在「新元挂牌利率」，不保留该外币挂牌分 sheet。
    _remove_sheet_if_exists(wb, "挂牌_SGD")

    wb.save(temp_workbook_path)

    # -----------------------------------------------------------------------
    # 2) overlay 写入数据区：设置 startrow
    # -----------------------------------------------------------------------
    # startrow 取值逻辑（pandas/openpyxl 的 0-based row 语义）：
    # - 元数据/促销类：表头在 Excel 第 1 行，数据从第 2 行开始 => startrow=1
    # - FX board：标题在第 1 行（合并），表头在第 2 行，数据从第 3 行开始 => startrow=2
    # 0-based startrow:
    # - header row is at Excel row 1 => startrow=0 when writing header
    # - data starts at Excel row 2 => header=False startrow=1
    data_startrow_1 = 1
    # - FX board: title row 1 (merged), header row 2 => data starts at Excel row 3
    data_startrow_2 = 2

    with pd.ExcelWriter(
        temp_workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="overlay",
    ) as writer:
        # 2.1) 元数据等：仅写数据区（不覆盖模板表头）
        df_meta.to_excel(
            writer, sheet_name="元数据", index=False, header=False, startrow=data_startrow_1
        )
        df_sgd.to_excel(
            writer,
            sheet_name="新元定存促销",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_sgd_board.to_excel(
            writer,
            sheet_name="新元挂牌利率",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_fx.to_excel(
            writer,
            sheet_name="外币定存促销",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_fx_notes.to_excel(
            writer,
            sheet_name="外币定存备注",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_other.to_excel(
            writer,
            sheet_name="其它利率",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )

        # 2.2) 外币挂牌按币种 sheet：仅写数据区（模板标题/列名保持）
        if fx_groups_sorted:
            for cur_key, sub_df in fx_groups_sorted:
                label = _fx_board_currency_label(cur_key)
                sn_fb = f"挂牌_{label}"
                sub_df.to_excel(
                    writer,
                    sheet_name=sn_fb,
                    index=False,
                    header=False,
                    startrow=data_startrow_2,
                )

        # -------------------------------------------------------------------
        # 3) 应用样式：按表类型选择 header_row
        # -------------------------------------------------------------------
        for sn in ("新元定存促销", "外币定存促销", "新元挂牌利率"):
            ws = writer.book[sn]
            _apply_bank_colors(ws, source_col_name="数据来源", header_row=1)
            _center_align_rate_columns(ws, header_row=1)
            if sn == "新元挂牌利率":
                _style_sgd_board_sheet(ws)

        for sn in fx_board_sheet_names:
            ws = writer.book[sn]
            _apply_bank_colors(ws, source_col_name="数据来源", header_row=2)
            _center_align_rate_columns(ws, header_row=2)

    # -----------------------------------------------------------------------
    # 4) 输出到最终路径
    # -----------------------------------------------------------------------
    shutil.copyfile(temp_workbook_path, path)


# Safety net: 自动给本模块中所有 def 加异常兜底
from bank_safety import auto_wrap_module_functions as _auto_wrap_module_functions

_auto_wrap_module_functions(globals(), module_name=__name__)

