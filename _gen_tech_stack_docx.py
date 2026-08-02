# -*- coding: utf-8 -*-
"""生成 RateStats 技术栈清单 Word（便于打印）。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "RateStats_技术栈清单.docx"


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    def add_title(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run(text), size=18, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))

    def add_subtitle(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        set_run_font(p.add_run(text), size=10, color=RGBColor(0x55, 0x55, 0x55))

    def add_h1(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(text), size=13, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))

    def add_bullet(text: str, bold_prefix: str | None = None) -> None:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.5)
        if bold_prefix:
            set_run_font(p.add_run(bold_prefix), size=10.5, bold=True)
            set_run_font(p.add_run(text), size=10.5)
        else:
            set_run_font(p.add_run(text), size=10.5)

    def add_body(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run(text), size=10.5)

    def add_table(headers: list[str], rows: list[list[str]]) -> None:
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            set_run_font(cell.paragraphs[0].add_run(h), size=10, bold=True)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = table.rows[ri + 1].cells[ci]
                cell.text = ""
                set_run_font(cell.paragraphs[0].add_run(val), size=9.5)
        doc.add_paragraph()

    add_title("RateStats 技术栈清单")
    add_subtitle("项目技术一览 · 便于打印归档")

    add_body(
        "RateStats 是一套新加坡银行存款利率采集与交付流水线："
        "AI 发现候选 URL → 规则+ML 打分选链 → HTML/PDF/API 解析 → "
        "输出 MarketRateData 与彩虹表 Excel。核心分三大模块。"
    )

    add_h1("一、三大模块")
    add_table(
        ["模块", "职责", "主要技术焦点"],
        [
            [
                "RateStats_Portable",
                "抓取解析、Excel、邮件、定时、Vertex 发现",
                "requests / BS4 / openpyxl / Vertex / Task Scheduler",
            ],
            [
                "RateStats_ML",
                "URL 排序模型训练、评估、调参、ML 选链",
                "scikit-learn / joblib / numpy",
            ],
            [
                "AI_Compare",
                "多搜索引擎评测与对比",
                "Vertex / Serper / Brave / Tavily",
            ],
        ],
    )

    add_h1("二、语言与运行环境")
    add_bullet("主语言；Windows 建议 3.10+，Mac 文档偏好 3.11/3.12", bold_prefix="Python — ")
    add_bullet("一键流水线入口（Portable / ML / AI_Compare）", bold_prefix="Windows Batch (.bat) — ")
    add_bullet("Mac 安装与流水线入口", bold_prefix="Bash / .command — ")
    add_bullet("一次性计划任务等辅助脚本", bold_prefix="PowerShell (.ps1) — ")

    add_h1("三、数据与 Excel")
    add_bullet("表格处理、评测报告", bold_prefix="pandas — ")
    add_bullet("写入 MarketRateData、彩虹表及各类报告", bold_prefix="openpyxl — ")
    add_bullet("配置、缓存、利率页、巡检报告、运行日志", bold_prefix="xlsx / JSON / CSV / PDF / HTML / .log — ")
    add_bullet("按日期归档流水线产物", bold_prefix="runs/YYYYMMDD/ — ")

    add_h1("四、抓取与解析")
    add_bullet("HTTP 主客户端", bold_prefix="requests — ")
    add_bullet("难抓站点的 TLS 指纹伪装", bold_prefix="curl_cffi — ")
    add_bullet("银行网页 HTML 解析", bold_prefix="BeautifulSoup4 + lxml — ")
    add_bullet("RHB / HSBC 等 PDF 利率表", bold_prefix="pypdf / pdfplumber — ")
    add_bullet("期限、利率等文本抽取", bold_prefix="正则表达式 — ")
    add_body("说明：selenium / webdriver-manager 出现在依赖声明中，当前业务代码基本未使用。")

    add_h1("五、AI 搜索与外部服务")
    add_bullet("生产侧 URL 发现（Discovery Engine）", bold_prefix="Google Vertex AI Search — ")
    add_bullet("多引擎评测对比", bold_prefix="Serper / Brave / Tavily — ")
    add_bullet("有接入但已标废弃", bold_prefix="Bing Web Search — ")
    add_bullet("Vertex 服务账号鉴权", bold_prefix="google-auth — ")
    add_bullet("报告邮件发送（配置中心可改 SMTP）", bold_prefix="SMTP — ")
    add_bullet("DBS / BEA / UOB 等公开网页、PDF、JSON API", bold_prefix="银行公开接口 — ")

    add_h1("六、机器学习（选链）")
    add_bullet("HistGradientBoostingClassifier，URL 排序", bold_prefix="scikit-learn — ")
    add_bullet("模型落盘（url_ranker.joblib）", bold_prefix="joblib — ")
    add_bullet("特征矩阵", bold_prefix="numpy — ")
    add_bullet("规则打分 + ML 概率融合；域名白名单硬过滤", bold_prefix="Rules + ML — ")

    add_h1("七、界面与运维部署")
    add_bullet("「配置中心」：邮件 + 定时任务；API Key 设置界面", bold_prefix="tkinter — ")
    add_bullet("schtasks 注册周任务（如每周三 17:55）", bold_prefix="Windows Task Scheduler — ")
    add_bullet("LaunchAgents 定时运行", bold_prefix="macOS launchd — ")
    add_bullet("Windows / Mac 便携分发包", bold_prefix="ZIP 打包 — ")

    add_h1("八、文档与演示（辅助）")
    add_bullet("演示文稿 / 海报生成", bold_prefix="python-pptx — ")
    add_bullet("周报、IA 报告、流程说明", bold_prefix="python-docx — ")
    add_bullet("ML 工作流示意图", bold_prefix="matplotlib — ")

    add_h1("九、端到端流水线（对应实现技术）")
    add_table(
        ["步骤", "做什么", "关键技术"],
        [
            ["1. Discover", "AI 搜索候选 URL", "Vertex AI Search / Discovery Engine"],
            ["2. Score", "规则 + ML 选链、域名白名单", "url_host_rules + sklearn picker"],
            ["3. Extract", "HTML / PDF / API 解析", "BS4 / pypdf / pdfplumber / JSON API"],
            ["4. Output", "MarketRateData + 彩虹表", "pandas / openpyxl"],
            ["Testing", "对照黄金 URL、多引擎评测、健康检查", "picker_eval / AI_Compare / check_url_health"],
            ["Deployment", "批处理 + 计划任务 → runs/日期", "bat + Task Scheduler / launchd + tkinter"],
        ],
    )

    add_h1("十、精简一页版（可贴幻灯片）")
    add_body(
        "Python · pandas / openpyxl · BeautifulSoup / PDF · "
        "Vertex + 多引擎搜索 · scikit-learn · Windows Task Scheduler · tkinter"
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    set_run_font(
        p.add_run("文档由仓库依赖与源码用法整理，仅列项目中实际使用的技术。"),
        size=9,
        color=RGBColor(0x77, 0x77, 0x77),
    )

    doc.save(OUT)
    print(f"[OK] {OUT}")


if __name__ == "__main__":
    main()
