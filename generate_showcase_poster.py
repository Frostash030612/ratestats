#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""NUS-ISS Project Showcase Poster (single-slide ppt) for RateStats."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "docs" / "IA_Reports"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x0B, 0x2C, 0x5E)
GOLD = RGBColor(0xB8, 0x8A, 0x2E)
GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x1A, 0x6B, 0x5C)


def _run(p, text, *, size=12, bold=False, color=GRAY):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return r


def _box(slide, left, top, width, height, fill=LIGHT, line=NAVY):
    sh = slide.shapes.add_shape(1, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1)
    return sh


def _textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def build() -> Path:
    prs = Presentation()
    # Widescreen poster-like single slide (same as presentation size; content dense)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # White background
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()

    # Top header bar
    header = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.95))
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()
    gold = slide.shapes.add_shape(1, 0, Inches(0.95), prs.slide_width, Inches(0.06))
    gold.fill.solid()
    gold.fill.fore_color.rgb = GOLD
    gold.line.fill.background()

    # Course name (do not change per template)
    t = _textbox(slide, Inches(0.35), Inches(0.12), Inches(9.5), Inches(0.35))
    p = t.text_frame.paragraphs[0]
    _run(p, "Graduate Diploma in Systems Analysis", size=14, bold=True, color=WHITE)

    # Project title
    t2 = _textbox(slide, Inches(0.35), Inches(0.42), Inches(10.5), Inches(0.45))
    p2 = t2.text_frame.paragraphs[0]
    _run(
        p2,
        "RateStats: From Manual Rate Collection to Python Automation + AI Search",
        size=20,
        bold=True,
        color=WHITE,
    )

    # Members
    t3 = _textbox(slide, Inches(0.35), Inches(0.72), Inches(8), Inches(0.25))
    p3 = t3.text_frame.paragraphs[0]
    _run(p3, "Group member: Xu Wenzhe  |  Host: Bank of China Singapore Branch (FinTech)", size=11, color=RGBColor(0xD0, 0xD8, 0xE8))

    # Logo placeholders (top right)
    logo = _box(slide, Inches(11.0), Inches(0.12), Inches(2.0), Inches(0.72), fill=RGBColor(0x15, 0x3A, 0x6E), line=GOLD)
    lt = _textbox(slide, Inches(11.05), Inches(0.25), Inches(1.9), Inches(0.5))
    lp = lt.text_frame.paragraphs[0]
    lp.alignment = PP_ALIGN.CENTER
    _run(lp, "[NUS-ISS Logo]", size=10, color=WHITE)

    # ---- Left: Introduction and Objective ----
    _box(slide, Inches(0.3), Inches(1.2), Inches(4.0), Inches(2.85))
    h = _textbox(slide, Inches(0.45), Inches(1.3), Inches(3.7), Inches(0.35))
    _run(h.text_frame.paragraphs[0], "Introduction and Objective", size=14, bold=True, color=NAVY)

    body = _textbox(slide, Inches(0.45), Inches(1.65), Inches(3.7), Inches(2.25))
    tf = body.text_frame
    tf.word_wrap = True
    lines = [
        ("Introduction", True),
        ("Banks need Singapore market deposit/promo rates across ~17 banks (~47 nodes). Work began as fully manual copy-paste into Excel.", False),
        ("Objective", True),
        ("1) Build a Python app to fetch/parse rates and output MarketRateData + Rainbow Excel.", False),
        ("2) Then plug AI Search into the same pipeline to discover & validate URLs — controlled AI, not cloud LLM agents.", False),
    ]
    first = True
    for text, is_h in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(4)
        _run(p, text, size=11 if is_h else 10, bold=is_h, color=NAVY if is_h else GRAY)

    # ---- Center: Pipeline diagram ----
    _box(slide, Inches(4.5), Inches(1.2), Inches(5.0), Inches(2.85))
    h = _textbox(slide, Inches(4.65), Inches(1.3), Inches(4.7), Inches(0.3))
    _run(h.text_frame.paragraphs[0], "Pipeline (build order)", size=14, bold=True, color=NAVY)

    stages = [
        ("0 Manual", "Copy rates\nby hand"),
        ("1 Python*", "Fetch/Parse\n→ Excel"),
        ("2 AI Search", "Find URLs"),
        ("3 Rules+ML", "Validate"),
    ]
    sw = Inches(1.05)
    sx = Inches(4.65)
    sy = Inches(1.75)
    for i, (title, body_t) in enumerate(stages):
        x = sx + i * Inches(1.2)
        box = _box(slide, x, sy, sw, Inches(1.35), fill=WHITE, line=TEAL if i >= 1 else RGBColor(0x99, 0x66, 0x66))
        tb = _textbox(slide, x + Inches(0.05), sy + Inches(0.15), sw - Inches(0.1), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, title, size=10, bold=True, color=NAVY)
        for line in body_t.split("\n"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            _run(p2, line, size=9, color=GRAY)
        if i < 3:
            ar = _textbox(slide, x + sw - Inches(0.05), sy + Inches(0.5), Inches(0.25), Inches(0.35))
            ap = ar.text_frame.paragraphs[0]
            ap.alignment = PP_ALIGN.CENTER
            _run(ap, "→", size=14, bold=True, color=GOLD)

    note = _textbox(slide, Inches(4.65), Inches(3.25), Inches(4.7), Inches(0.65))
    nf = note.text_frame
    nf.word_wrap = True
    p = nf.paragraphs[0]
    _run(p, "* Steps 1 (Python extract/Excel) built first; AI Search + scoring added later into the same flow. Output: MarketRateData & Rainbow Table.", size=9, color=GRAY)

    # ---- Right: Solutions ----
    _box(slide, Inches(9.7), Inches(1.2), Inches(3.35), Inches(2.85))
    h = _textbox(slide, Inches(9.85), Inches(1.3), Inches(3.05), Inches(0.35))
    _run(h.text_frame.paragraphs[0], "Solutions / Key Features", size=13, bold=True, color=NAVY)
    body = _textbox(slide, Inches(9.85), Inches(1.7), Inches(3.05), Inches(2.2))
    tf = body.text_frame
    tf.word_wrap = True
    sols = [
        "Python extractors: HTML / PDF / API (HSBC, DBS, OCBC…)",
        "AI Search (Vertex) proposes candidate URLs",
        "Hybrid filter: host whitelist + ML ranker (~95% hit)",
        "Batch + Task Scheduler → dated runs/",
        "Multi-AI benchmark: Vertex / Serper / Brave / Tavily",
    ]
    for i, s in enumerate(sols):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        _run(p, f"• {s}", size=10, color=GRAY)

    # ---- Bottom left: Challenges ----
    _box(slide, Inches(0.3), Inches(4.2), Inches(6.4), Inches(2.55))
    h = _textbox(slide, Inches(0.45), Inches(4.3), Inches(6.1), Inches(0.3))
    _run(h.text_frame.paragraphs[0], "Challenges, Lessons Learnt", size=14, bold=True, color=NAVY)
    body = _textbox(slide, Inches(0.45), Inches(4.65), Inches(6.1), Inches(1.95))
    tf = body.text_frame
    tf.word_wrap = True
    challenges = [
        "Challenge: AI returned wrong-bank pages (e.g. ICBC → HLF) that looked plausible.",
        "Lesson: Domain whitelist before ML; hybrid Rules + ML beats pure AI ranking.",
        "Challenge: Same bank, wrong page type (OCBC Premier vs SGD board).",
        "Lesson: Hard-reject patterns + hard_negatives.csv to teach decision boundaries.",
        "Lesson: Build a usable Python product first; plug AI into a real pipeline, not a demo.",
        "Lesson: Banking security favours controlled AI Search over unrestricted cloud agents.",
    ]
    for i, s in enumerate(challenges):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        _run(p, f"• {s}", size=10, color=GRAY)

    # ---- Bottom right: Impact + company ----
    _box(slide, Inches(6.9), Inches(4.2), Inches(6.15), Inches(2.55))
    h = _textbox(slide, Inches(7.05), Inches(4.3), Inches(5.85), Inches(0.3))
    _run(h.text_frame.paragraphs[0], "Business Impact & Recommendations", size=14, bold=True, color=NAVY)
    body = _textbox(slide, Inches(7.05), Inches(4.65), Inches(5.85), Inches(1.5))
    tf = body.text_frame
    tf.word_wrap = True
    impact = [
        "Impact: Manual copy → Python Excel app → less URL hunting with AI (~95% hit).",
        "Impact: Schedulable MarketRateData / Rainbow for market comparison.",
        "Recommend: Keep expanding hard_negatives; re-tune picker after each batch.",
        "Recommend: Scale controlled AI discovery to other external-data monitoring.",
    ]
    for i, s in enumerate(impact):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        _run(p, f"• {s}", size=10, color=GRAY)

    # Company logo placeholder
    cl = _box(slide, Inches(10.5), Inches(6.15), Inches(2.3), Inches(0.45), fill=WHITE, line=GOLD)
    ct = _textbox(slide, Inches(10.55), Inches(6.22), Inches(2.2), Inches(0.35))
    cp = ct.text_frame.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    _run(cp, "[BOC Logo]", size=10, bold=True, color=NAVY)

    # Footer
    ft = _textbox(slide, Inches(0.3), Inches(6.9), Inches(10), Inches(0.35))
    fp = ft.text_frame.paragraphs[0]
    _run(
        fp,
        "Industrial Attachment  ·  Bank of China Singapore  ·  Xu Wenzhe  ·  Project Showcase Poster (ppt)",
        size=9,
        color=RGBColor(0x88, 0x88, 0x88),
    )

    path = OUT / "Project_Showcase_Poster_Xu_Wenzhe.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build())
