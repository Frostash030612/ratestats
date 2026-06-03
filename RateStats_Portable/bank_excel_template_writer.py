from __future__ import annotations

"""
可移植版本的 MarketRateData 模板写入器。

该文件与主目录 `bank_excel_template_writer.py` 逻辑一致，仅用于当
你运行 `RateStats_Portable/bank_all_promo_rates.py` 时，避免依赖主目录
模块解析路径。
"""

import os
import shutil
from datetime import datetime
from typing import Any, List, Tuple

import pandas as pd

from bank_excel_style_helpers import (
    apply_bank_colors as _apply_bank_colors,
    center_align_rate_columns as _center_align_rate_columns,
    style_sgd_board_sheet as _style_sgd_board_sheet,
)

from bank_safety import safe_call


def _is_empty_rate_cell(v: Any) -> bool:
    "判断利率单元格是否为空值（含 N/A 文本）。"
    if v is None:
        return True
    t = str(v).strip().upper()
    return t == "" or t == "N/A"


def _trim_fx_board_trailing_empty_tenors(
    ws: Any,
    *,
    header_row: int,
    data_start_row: int,
) -> None:
    """
    对外币挂牌 sheet 裁剪“尾部全空期限列”。

    例如最后一个有值的期限是 18M，则删除其后的 24M/36M/48M/60M 列，
    避免表格右侧出现大量无意义空列。
    """
    tenor_order = [
        "7D_pct",
        "14D_pct",
        "1M_pct",
        "2M_pct",
        "3M_pct",
        "6M_pct",
        "9M_pct",
        "12M_pct",
        "18M_pct",
        "24M_pct",
        "36M_pct",
        "48M_pct",
        "60M_pct",
    ]
    header_to_col: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        hv = ws.cell(row=header_row, column=c).value
        if hv is None:
            continue
        key = str(hv).strip()
        if key in tenor_order:
            header_to_col[key] = c

    present_tenors = [t for t in tenor_order if t in header_to_col]
    if not present_tenors:
        return

    if ws.max_row < data_start_row:
        return

    last_non_empty_idx = -1
    for i, tenor in enumerate(present_tenors):
        col = header_to_col[tenor]
        has_value = False
        for r in range(data_start_row, ws.max_row + 1):
            if not _is_empty_rate_cell(ws.cell(row=r, column=col).value):
                has_value = True
                break
        if has_value:
            last_non_empty_idx = i

    # 全空则不裁剪，保持模板原样。
    if last_non_empty_idx < 0:
        return

    to_delete = present_tenors[last_non_empty_idx + 1 :]
    for tenor in reversed(to_delete):
        ws.delete_cols(header_to_col[tenor], 1)


def _fx_board_currency_label(cur: Any) -> str:
    """
    将 payload 里的“币种值”归一为用于 sheet 标题的短标签。

    主要作用：
    1) 避免 None / 空字符串导致 sheet 名称不稳定；
    2) 将 NaN 视为“未填币种”，统一落到一个可预期的标题后缀。
    """
    if cur is None:
        return "（未填币种）"
    if isinstance(cur, float):
        try:
            if pd.isna(cur):
                return "（未填币种）"
        except Exception:
            pass
    t = str(cur).strip()
    return t if t else "（未填币种）"


def _is_sgd_fx_board_currency(cur: Any) -> bool:
    """SGD 挂牌走「新元挂牌利率」sheet，不生成外币挂牌分 sheet。"""
    return _fx_board_currency_label(cur).strip().upper() == "SGD"


def _remove_sheet_if_exists(wb: Any, name: str) -> None:
    "从工作簿删除指定工作表；若不存在则忽略。"
    if name in wb.sheetnames:
        wb.remove(wb[name])


@safe_call(
    context="写入 MarketRateData Excel（portable 模板覆写）",
    default_value=None,
)
def write_market_rate_excel_with_template(
    *,
    df_meta: pd.DataFrame,
    df_sgd: pd.DataFrame,
    df_sgd_board: pd.DataFrame,
    df_fx: pd.DataFrame,
    df_fx_notes: pd.DataFrame,
    df_other: pd.DataFrame,
    df_fx_board: pd.DataFrame,
    path: str,
    script_dir: str,
) -> None:
    """
    使用 `assets/MarketRateData_template.xlsx` 作为固定模板，将传入的 DataFrame
    覆盖/写入到正确的 sheet 与区域（不覆盖模板表头）。

    详见主目录同名模块：`bank_excel_template_writer.py`。
    """

    # 注意：这里不直接在原模板文件上写，而是 copy 到临时文件，避免破坏版本化模板。
    template_path_candidates = [
        os.path.join(script_dir, "assets", "MarketRateData_template.xlsx"),
        os.path.join(script_dir, "..", "assets", "MarketRateData_template.xlsx"),
    ]
    template_path = next(
        (p for p in template_path_candidates if os.path.exists(p)), None
    )
    if not template_path:
        raise FileNotFoundError(
            "Cannot find MarketRateData_template.xlsx in assets/. "
            f"Looked in: {template_path_candidates}"
        )

    temp_dir = os.path.join(script_dir, "temp")
    os.makedirs(temp_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    temp_workbook_path = os.path.join(
        temp_dir, f"MarketRateData_work_{ts}.xlsx"
    )
    shutil.copyfile(template_path, temp_workbook_path)

    from openpyxl import load_workbook

    wb = load_workbook(temp_workbook_path)

    fx_board_sheet_names: List[str] = []
    fx_base_sheet = "挂牌_USD"
    fx_empty_sheet = "外币挂牌利率"
    curr_col = "币种"

    fx_groups_sorted: List[Tuple[Any, pd.DataFrame]] = []
    if df_fx_board.empty or curr_col not in df_fx_board.columns:
        _remove_sheet_if_exists(wb, fx_empty_sheet)
    else:
        groups = list(df_fx_board.groupby(curr_col, dropna=False, sort=False))
        groups = [
            (k, v) for k, v in groups if not _is_sgd_fx_board_currency(k)
        ]
        groups.sort(key=lambda kv: _fx_board_currency_label(kv[0]))
        fx_groups_sorted = groups

        if not fx_groups_sorted:
            _remove_sheet_if_exists(wb, fx_empty_sheet)
        else:
            for cur_key, sub_df in fx_groups_sorted:
                label = _fx_board_currency_label(cur_key)
                sn_fb = f"挂牌_{label}"
                if sn_fb not in wb.sheetnames:
                    if fx_base_sheet not in wb.sheetnames:
                        raise RuntimeError(
                            f"Template missing base fx board sheet {fx_base_sheet!r} "
                            f"needed to copy for new currency {label!r}."
                        )
                    base_ws = wb[fx_base_sheet]
                    new_ws = wb.copy_worksheet(base_ws)
                    new_ws.title = sn_fb

                wb[sn_fb]["A1"].value = f"{label} 外币挂牌利率"
                fx_board_sheet_names.append(sn_fb)

            _remove_sheet_if_exists(wb, fx_empty_sheet)

    _remove_sheet_if_exists(wb, "挂牌_SGD")
    # “其它利率”业务已下线：无论模板是否存在该 sheet，统一移除。
    _remove_sheet_if_exists(wb, "其它利率")

    wb.save(temp_workbook_path)

    data_startrow_1 = 1
    data_startrow_2 = 2

    with pd.ExcelWriter(
        temp_workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="overlay",
    ) as writer:
        df_meta.to_excel(
            writer, sheet_name="元数据", index=False, header=False, startrow=data_startrow_1
        )
        df_sgd.to_excel(
            writer,
            sheet_name="新元定存促销",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_sgd_board.to_excel(
            writer,
            sheet_name="新元挂牌利率",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_fx.to_excel(
            writer,
            sheet_name="外币定存促销",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        df_fx_notes.to_excel(
            writer,
            sheet_name="外币定存备注",
            index=False,
            header=False,
            startrow=data_startrow_1,
        )
        if fx_groups_sorted:
            for cur_key, sub_df in fx_groups_sorted:
                label = _fx_board_currency_label(cur_key)
                sn_fb = f"挂牌_{label}"
                sub_df.to_excel(
                    writer,
                    sheet_name=sn_fb,
                    index=False,
                    header=False,
                    startrow=data_startrow_2,
                )

        for sn in ("新元定存促销", "外币定存促销", "新元挂牌利率"):
            ws = writer.book[sn]
            _apply_bank_colors(ws, source_col_name="数据来源", header_row=1)
            _center_align_rate_columns(ws, header_row=1)
            if sn == "新元挂牌利率":
                _style_sgd_board_sheet(ws)

        for sn in fx_board_sheet_names:
            ws = writer.book[sn]
            _trim_fx_board_trailing_empty_tenors(
                ws,
                header_row=2,
                data_start_row=3,
            )
            _apply_bank_colors(ws, source_col_name="数据来源", header_row=2)
            _center_align_rate_columns(ws, header_row=2)

    shutil.copyfile(temp_workbook_path, path)


# Safety net: 自动给本模块中所有 def 加异常兜底
from bank_safety import auto_wrap_module_functions as _auto_wrap_module_functions

_auto_wrap_module_functions(globals(), module_name=__name__)

