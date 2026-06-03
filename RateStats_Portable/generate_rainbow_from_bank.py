#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
基于 MarketRateData 输出表，按“彩虹表 20260311.xlsx”样式输出新彩虹表。

说明：
1) 样式来自模板（不会把模板中的旧数据当来源）
2) 数据来源来自 source Excel（默认最新 MarketRateData_*.xlsx）
3) 当前更新的模板 sheet：
   - SGD Promotional Rate
   - SGD Board Rate
   - SGD Board Rate Ranked
   - Sheet3（SGD board 矩阵）
"""

from __future__ import annotations

import argparse
import datetime as dt
import glob
import math
import os
import re
import shutil
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


RATE_THRESHOLDS = [
    ("red", 0.0140),
    ("orange", 0.0130),
    ("yellow", 0.0120),
    ("green", 0.0110),
    ("blue", 0.0080),
]

FILL_BY_NAME = {
    # Lighter pastel palette for rainbow table.
    "red": PatternFill("solid", fgColor="FFF4CCCC"),
    "orange": PatternFill("solid", fgColor="FFFCE5CD"),
    "yellow": PatternFill("solid", fgColor="FFFFF2CC"),
    "green": PatternFill("solid", fgColor="FFD9EAD3"),
    "blue": PatternFill("solid", fgColor="FFD0E0E3"),
    "purple": PatternFill("solid", fgColor="FFD9D2E9"),
}

TENORS_PROMO = ["1M", "3M", "5M", "6M", "9M", "12M"]
TENORS_BOARD = ["7D", "1M", "2M", "3M", "4M", "5M", "6M", "7M", "8M", "9M", "10M", "11M", "12M", "18M", "24M", "36M"]


def normalize_text(s: Any) -> str:
    "统一清洗文本，便于做银行名与金额区间匹配。"
    t = str(s or "").strip().lower()
    t = t.replace(" ", "").replace("\u3000", "").replace(",", "")
    return t


def extract_number_tokens(s: Any) -> set[str]:
    "从金额文本中提取数字 token（支持 k/m/mil）。"
    txt = str(s or "")
    out: set[str] = set()
    for m in re.finditer(r"\d+(?:,\d+)*(?:\s*(?:k|m|mil))?", txt, flags=re.I):
        tok = m.group(0).lower().replace(",", "").strip().replace("mil", "m")
        if tok.endswith("k") and tok[:-1].isdigit():
            out.add(str(int(tok[:-1]) * 1000))
            continue
        if tok.endswith("m") and tok[:-1].isdigit():
            out.add(str(int(tok[:-1]) * 1000000))
            continue
        if tok.isdigit():
            out.add(tok)
    return out


def score_amount_match(template_amount: Any, source_amount: Any) -> int:
    "为模板金额与来源金额打分，选择最匹配记录。"
    t = str(template_amount or "")
    s = str(source_amount or "")
    if not t.strip() or not s.strip():
        return 0
    t_norm = normalize_text(t)
    s_norm = normalize_text(s)
    if t_norm == s_norm:
        return 40
    score = 0
    t_tokens = extract_number_tokens(t)
    s_tokens = extract_number_tokens(s)
    inter = t_tokens & s_tokens
    if inter:
        score += 10 * len(inter)
    # Directional cue matching: < / <= / up to / below / 以下  vs  > / >= / above / 及以上
    t_low = any(k in t_norm for k in ["<", "<=", "≤", "＜", "upto", "below", "lessthan", "以下"])
    t_high = any(k in t_norm for k in [">", ">=", "≥", "＞", "above", "atleast", "andabove", "及以上"])
    s_low = any(k in s_norm for k in ["<", "<=", "≤", "＜", "upto", "below", "lessthan", "以下"])
    s_high = any(k in s_norm for k in [">", ">=", "≥", "＞", "above", "atleast", "andabove", "及以上"])
    if t_low and s_low:
        score += 12
    if t_high and s_high:
        score += 12
    # Range-like cue
    t_range = any(k in t_norm for k in ["-", "to", "至"])
    s_range = any(k in s_norm for k in ["-", "to", "至"])
    if t_range and s_range:
        score += 6
    return score


def rate_to_color(rate_decimal: float) -> str:
    "按利率阈值返回彩虹色名称。"
    for name, thr in RATE_THRESHOLDS:
        if rate_decimal >= thr:
            return name
    return "purple"


def pct_to_decimal(v: Any) -> Optional[float]:
    "把百分比值安全转换为小数利率。"
    if v is None:
        return None
    if isinstance(v, str) and not v.strip():
        return None
    try:
        fv = float(v)
    except (TypeError, ValueError):
        return None
    if math.isnan(fv):
        return None
    return fv / 100.0


def maybe_parse_decimal(v: Any) -> Optional[float]:
    "宽松解析数值文本为小数，失败返回 None。"
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def canonical_bank_name(s: Any) -> str:
    "归一化银行名称，兼容中英文别名。"
    raw = str(s or "").strip()
    up = raw.upper()
    alias = {
        "BOC SG": "BOC",
        "SBI SG": "SBI",
        "HL BANK": "HL BANK",
        "SINGFINANCE": "SINGFINANCE",
    }
    for k, v in alias.items():
        if k in up:
            return v
    return up


def get_effective_cell_value(ws, row: int, col: int) -> Any:
    "读取合并单元格中的有效值（向上追溯）。"
    v = ws.cell(row, col).value
    if v is not None:
        return v
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = rng.bounds
        if min_col <= col <= max_col and min_row <= row <= max_row:
            return ws.cell(min_row, min_col).value
    return None


def detect_template_sgd_promo_tenors(ws) -> Dict[str, int]:
    "识别 SGD Promotional Rate 表头中的期数字段。"
    out: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(4, c).value
        if isinstance(v, str) and v.strip() in {"1M", "3M", "6M", "9M", "12M"}:
            out[v.strip()] = c
    return out


def detect_template_sgd_board_tenor_blocks(ws, header_row: int) -> Dict[str, tuple[int, int, int]]:
    "识别 SGD Board 表中的利率列分组信息。"
    out: Dict[str, tuple[int, int, int]] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and v.strip() in {"7D", "1M", "3M", "6M", "9M", "12M", "18M", "24M", "36M"}:
            out[v.strip()] = (c, c + 1, c + 2)  # bank/rate/amount
    return out


def apply_colors_to_rate_columns(ws, rate_cols: List[int]) -> None:
    "根据利率区间给目标列批量着色。"
    for r in range(1, ws.max_row + 1):
        for c in rate_cols:
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                ws.cell(r, c).fill = FILL_BY_NAME[rate_to_color(float(v))]


def pick_best_promo_row(df_bank: pd.DataFrame, template_cond: Any) -> Optional[pd.Series]:
    "在促销数据中挑选与模板行最匹配的一条记录。"
    if df_bank.empty:
        return None
    cond_text = str(template_cond or "")
    cond_norm = normalize_text(cond_text)
    cond_tokens = extract_number_tokens(cond_text)
    best_idx = None
    best_score = -1
    for i, row in df_bank.iterrows():
        s = " ".join(
            str(row.get(k, "") or "")
            for k in ("产品或客群", "资金区间_页面", "起存金额_页面", "上限金额_页面")
        )
        n = normalize_text(s)
        score = 0
        if cond_norm and cond_norm in n:
            score += 10
        score += 5 * len(cond_tokens & extract_number_tokens(s))
        if score > best_score:
            best_score = score
            best_idx = i
    if best_idx is None:
        return None
    return df_bank.loc[best_idx]


def find_best_board_row(df_board: pd.DataFrame, bank_name: str, template_amount: Any) -> Optional[pd.Series]:
    "在挂牌数据中按银行+金额区间寻找最佳匹配行。"
    cbank = canonical_bank_name(bank_name)
    d = df_board[df_board["_bank_canon"] == cbank]
    if d.empty:
        return None
    best_idx = None
    best_score = -1
    for i, row in d.iterrows():
        s = " ".join(
            str(row.get(k, "") or "")
            for k in ("资金区间_页面", "起存金额_页面", "上限金额_页面", "产品或档位")
        )
        sc = score_amount_match(template_amount, s)
        if sc > best_score:
            best_score = sc
            best_idx = i
    if best_idx is None:
        return None
    return d.loc[best_idx]


def _is_amount_like(v: Any) -> bool:
    "判断文本是否像“金额/资金区间”描述。"
    if v is None:
        return False
    if isinstance(v, (int, float)):
        # pure numeric values are usually rates, not amount text
        return False
    s = str(v).strip()
    if not s:
        return False
    s_norm = normalize_text(s)
    keys = ["$", "sgd", "usd", "<", ">", "≤", "≥", "＜", "＞", "-", "to", "至", "以下", "及以上", "k", "m", "mil"]
    return any(k in s_norm for k in keys)


def resolve_rate_amount_cols(ws, r: int, rate_col_guess: int, amount_col_guess: int) -> tuple[int, int]:
    "自动识别源数据中的利率列与金额区间列。"
    rv = get_effective_cell_value(ws, r, rate_col_guess)
    av = get_effective_cell_value(ws, r, amount_col_guess)
    r_is_rate = maybe_parse_decimal(rv) is not None
    a_is_rate = maybe_parse_decimal(av) is not None
    r_is_amt = _is_amount_like(rv)
    a_is_amt = _is_amount_like(av)
    # If looks reversed (amount in rate col, rate in amount col), swap.
    if r_is_amt and a_is_rate:
        return amount_col_guess, rate_col_guess
    # Normal case.
    if r_is_rate and a_is_amt:
        return rate_col_guess, amount_col_guess
    # Fallback to template layout.
    return rate_col_guess, amount_col_guess


def normalize_rate_amount_layout(ws, blocks: Dict[str, tuple[int, int, int]], start_row: int) -> None:
    """
    Fix template rows where rate/amount values are physically written in reversed columns.
    """
    for _, (_, rate_col, amount_col) in blocks.items():
        for r in range(start_row, ws.max_row + 1):
            rv = ws.cell(r, rate_col).value
            av = ws.cell(r, amount_col).value
            if _is_amount_like(rv) and (maybe_parse_decimal(av) is not None):
                ws.cell(r, rate_col).value = av
                ws.cell(r, amount_col).value = rv


def load_source_tables(source_path: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    "读取来源工作簿，提取促销与挂牌两类 DataFrame。"
    # 按约定的 sheet 名读取来源数据。
    df_promo = pd.read_excel(source_path, sheet_name="新元定存促销")
    df_board = pd.read_excel(source_path, sheet_name="新元挂牌利率")
    # 预先缓存银行名归一化字段，后续匹配时避免重复计算。
    df_promo["_bank_canon"] = df_promo["数据来源"].map(canonical_bank_name)
    df_board["_bank_canon"] = df_board["数据来源"].map(canonical_bank_name)
    return df_promo, df_board


def update_sgd_promotional_rate(ws, df_promo: pd.DataFrame) -> None:
    "把来源促销利率写入模板 SGD Promotional Rate。"
    tenor_label_cols = detect_template_sgd_promo_tenors(ws)
    # 按模板逐行遍历，用“银行 + 条件列”匹配来源行，再回填利率。
    for r in range(1, ws.max_row + 1):
        bank = ws.cell(r, 2).value
        if not isinstance(bank, str) or not bank.strip():
            continue
        bank_df = df_promo[df_promo["_bank_canon"] == canonical_bank_name(bank)]
        if bank_df.empty:
            continue
        for tenor, label_col in tenor_label_cols.items():
            rate_col = label_col + 1
            cond_col = label_col + 2
            picked = pick_best_promo_row(bank_df, ws.cell(r, cond_col).value)
            if picked is None:
                continue
            dec = pct_to_decimal(picked.get(f"{tenor}_pct"))
            if dec is not None:
                ws.cell(r, rate_col).value = float(dec)
    apply_colors_to_rate_columns(ws, [c + 1 for c in tenor_label_cols.values()])


def update_sgd_board_rate(ws, df_board: pd.DataFrame) -> None:
    "把来源挂牌利率写入模板 SGD Board Rate。"
    blocks = detect_template_sgd_board_tenor_blocks(ws, header_row=5)
    # 先修正模板中可能存在的“利率列/金额列”反置问题，再开始回填。
    normalize_rate_amount_layout(ws, blocks, start_row=6)
    for tenor, (bank_col, rate_col_guess, amount_col_guess) in blocks.items():
        src_col = f"{tenor}_pct"
        if src_col not in df_board.columns:
            continue
        current_bank: Optional[str] = None
        for r in range(6, ws.max_row + 1):
            bank_cell = get_effective_cell_value(ws, r, bank_col)
            if isinstance(bank_cell, str) and bank_cell.strip():
                current_bank = bank_cell.strip()
            bank = current_bank
            if not bank:
                continue
            _, amount_col = resolve_rate_amount_cols(ws, r, rate_col_guess, amount_col_guess)
            templ_amt = get_effective_cell_value(ws, r, amount_col)
            picked = find_best_board_row(df_board, bank, templ_amt)
            if picked is None:
                continue
            dec = pct_to_decimal(picked.get(src_col))
            if dec is None:
                continue
            rate_col, _ = resolve_rate_amount_cols(ws, r, rate_col_guess, amount_col_guess)
            ws.cell(r, rate_col).value = float(dec)
        # Color only Highest Rate cells within the actual block rows.
        current_bank_for_color: Optional[str] = None
        for r in range(6, ws.max_row + 1):
            b = get_effective_cell_value(ws, r, bank_col)
            if isinstance(b, str) and b.strip():
                current_bank_for_color = b.strip()
            if not current_bank_for_color:
                continue
            v = ws.cell(r, rate_col_guess).value
            if isinstance(v, (int, float)):
                ws.cell(r, rate_col_guess).fill = FILL_BY_NAME[rate_to_color(float(v))]


def update_sgd_board_ranked(ws, df_board: pd.DataFrame) -> None:
    "更新模板 SGD Board Rate Ranked 的利率区域。"
    blocks = detect_template_sgd_board_tenor_blocks(ws, header_row=1)
    normalize_rate_amount_layout(ws, blocks, start_row=2)
    for tenor, (bank_col, rate_col_guess, amount_col_guess) in blocks.items():
        src_col = f"{tenor}_pct"
        if src_col not in df_board.columns:
            continue
        # 对每个 tenor 先按利率降序得到排行榜，再按模板行回填。
        offers = df_board[["数据来源", "资金区间_页面", src_col]].copy()
        offers = offers.dropna(subset=[src_col])
        offers["_dec"] = offers[src_col].apply(pct_to_decimal)
        offers = offers.dropna(subset=["_dec"]).sort_values("_dec", ascending=False).reset_index(drop=True)
        rank_rows = [r for r in range(2, ws.max_row + 1) if isinstance(ws.cell(r, bank_col).value, str)]
        for i, r in enumerate(rank_rows):
            if i >= len(offers):
                break
            ws.cell(r, bank_col).value = str(offers.iloc[i]["数据来源"])
            rate_col, amount_col = resolve_rate_amount_cols(ws, r, rate_col_guess, amount_col_guess)
            ws.cell(r, rate_col).value = float(offers.iloc[i]["_dec"])
            ws.cell(r, amount_col).value = offers.iloc[i]["资金区间_页面"]
        # Color only Highest Rate cells for populated ranking rows.
        for r in rank_rows:
            v = ws.cell(r, rate_col_guess).value
            if isinstance(v, (int, float)):
                ws.cell(r, rate_col_guess).fill = FILL_BY_NAME[rate_to_color(float(v))]


def update_sheet3_matrix(ws, df_board: pd.DataFrame) -> None:
    "更新模板 Sheet3 的矩阵利率区域。"
    tenor_cols: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, str) and v.strip() in {"7D", "1M", "3M", "6M", "9M", "12M", "18M", "24M", "36M"}:
            tenor_cols[v.strip()] = c
    if not tenor_cols:
        return
    # 按银行与金额区间为矩阵每行匹配最佳来源记录。
    for r in range(1, ws.max_row + 1):
        bank = get_effective_cell_value(ws, r, 1)
        amt = get_effective_cell_value(ws, r, 2)
        if not isinstance(bank, str) or not bank.strip():
            continue
        picked = find_best_board_row(df_board, bank, amt)
        if picked is None:
            continue
        for tenor, c in tenor_cols.items():
            dec = pct_to_decimal(picked.get(f"{tenor}_pct"))
            if dec is not None:
                ws.cell(r, c).value = float(dec)
    apply_colors_to_rate_columns(ws, list(tenor_cols.values()))


def resolve_latest_source(work_dir: str) -> str:
    "自动选择目录中最新的可用来源工作簿。"
    pats = [
        "MarketRateData_*.xlsx",
        "marketratedata_*.xlsx",
        "bank_promo_rates_colored_centered_*.xlsx",
        "bank_promo_rates_colored_*.xlsx",
        "bank_promo_rates_*.xlsx",
    ]
    cands: List[str] = []
    for p in pats:
        cands.extend(glob.glob(os.path.join(work_dir, p)))
    if not cands:
        raise FileNotFoundError("Cannot find MarketRateData source workbook")
    cands.sort(key=lambda x: os.path.getmtime(x), reverse=True)
    return cands[0]


def resolve_template(work_dir: str, template_arg: str) -> str:
    "解析彩虹表模板路径（支持通配符）。"
    if "*" in template_arg or "?" in template_arg:
        cands = sorted(glob.glob(os.path.join(work_dir, template_arg)))
        if not cands:
            raise FileNotFoundError(f"No template matches: {template_arg}")
        return cands[0]
    p = os.path.join(work_dir, template_arg)
    if not os.path.exists(p):
        raise FileNotFoundError(p)
    return p


def _rainbow_zh_filename(run_tag: str) -> str:
    "与 03 批处理约定一致的中文输出名（仅含 ASCII 的 run_tag 由调用方传入）。"
    return f"彩虹表_按MarketRateData更新_{run_tag}.xlsx"


def main() -> None:
    "彩虹表生成入口：解析参数、复制模板并回填数据。"
    ap = argparse.ArgumentParser(description="从 MarketRateData 生成彩虹表（保留模板样式）。")
    ap.add_argument("--source", default=None, help="来源 Excel（默认自动取最新 MarketRateData*.xlsx）")
    ap.add_argument("--template", default="*20260311*.xlsx", help="彩虹表模板（支持通配符）")
    ap.add_argument("--out", default=None, help="输出 .xlsx 路径（相对脚本目录或绝对路径）")
    ap.add_argument(
        "--out-dir-zh",
        default=None,
        metavar="DIR",
        help="输出目录；与 --out-tag 同时使用时写入中文名彩虹表（避免经 cmd 传中文路径乱码）。",
    )
    ap.add_argument(
        "--out-tag",
        default=None,
        metavar="TAG",
        help="与 --out-dir-zh 搭配，例如 20260514_16.24",
    )
    args = ap.parse_args()

    work_dir = os.path.dirname(os.path.abspath(__file__))
    # 解析来源与模板路径（未显式传参时自动发现最新文件）。
    source_path = args.source if args.source else resolve_latest_source(work_dir)
    if not os.path.isabs(source_path):
        source_path = os.path.join(work_dir, source_path)
    if not os.path.exists(source_path):
        raise FileNotFoundError(source_path)
    template_path = resolve_template(work_dir, args.template)

    date_tag = dt.datetime.now().strftime("%Y%m%d_%H.%M")
    if args.out_dir_zh and args.out_tag:
        if args.out:
            ap.error("不要同时使用 --out 与 --out-dir-zh/--out-tag")
        out_path = os.path.join(os.path.abspath(args.out_dir_zh), _rainbow_zh_filename(args.out_tag))
    elif args.out_dir_zh or args.out_tag:
        ap.error("--out-dir-zh 与 --out-tag 须同时提供")
    elif args.out:
        out_path = os.path.join(work_dir, args.out)
    else:
        out_path = os.path.join(work_dir, f"彩虹表_从MarketRateData生成_{date_tag}.xlsx")
    if os.path.exists(out_path):
        base, ext = os.path.splitext(out_path)
        i = 1
        while os.path.exists(f"{base}_{i}{ext}"):
            i += 1
        out_path = f"{base}_{i}{ext}"

    # 先复制模板，再在副本上执行数据回填，保留模板样式。
    shutil.copyfile(template_path, out_path)

    df_promo, df_board = load_source_tables(source_path)
    wb = load_workbook(out_path)

    # 按存在的模板 sheet 分别执行更新，避免模板差异导致报错。
    if "SGD Promotional Rate" in wb.sheetnames:
        update_sgd_promotional_rate(wb["SGD Promotional Rate"], df_promo)
    if "SGD Board Rate" in wb.sheetnames:
        update_sgd_board_rate(wb["SGD Board Rate"], df_board)
    if "SGD Board Rate Ranked" in wb.sheetnames:
        update_sgd_board_ranked(wb["SGD Board Rate Ranked"], df_board)
    if "Sheet3" in wb.sheetnames:
        update_sheet3_matrix(wb["Sheet3"], df_board)

    wb.save(out_path)
    print(f"Generated: {out_path}")
    print(f"Source: {source_path}")
    print(f"Template: {template_path}")


if __name__ == "__main__":
    print("[INFO] 兼容入口：建议改用 generate_rainbow_from_market.py")
    main()
