#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 MarketRateData 生成「SGD Board Rate」工作表。

版式（与 SGD Promotional Rate 一致）：
- A1：日期 yyyymmdd
- 第 2 行起每 3 列一组：期限 / Highest Rate / Amount
- 期限顺序：7D, 1M, 3M, 6M, 9M, 12M, 18M, 24M, 36M（连续列，无空列）
- 每组浅色底色；同银行同利率的多个资金区间合并为一行
- 同银行多区间排在一起，银行名列纵向合并
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TENORS = ["7D", "1M", "3M", "6M", "9M", "12M", "18M", "24M", "36M"]

TENOR_START_COL = {tenor: 1 + idx * 3 for idx, tenor in enumerate(TENORS)}

TENOR_FILL_HEX = {
    "7D": "E3F2FD",
    "1M": "E8F4FD",
    "3M": "E8F5E9",
    "6M": "FFF8E1",
    "9M": "FCE4EC",
    "12M": "F3E5F5",
    "18M": "E0F2F1",
    "24M": "FFF3E0",
    "36M": "ECEFF1",
}

HEADER_ROW = 2
DATA_START_ROW = 3

thin = Side(border_style="thin", color="000000")
CELL_BORDER = Border(top=thin, left=thin, right=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_FONT = Font(bold=True)


def display_bank(name: object) -> str:
    s = str(name or "").strip()
    return re.sub(r"\s+SG$", "", s, flags=re.I).strip() or s


def date_tag_from_path(path: Path) -> str:
    m = re.search(r"MarketRateData_(\d{8})", path.name)
    if m:
        return m.group(1)
    m = re.search(r"(\d{8})", path.stem)
    return m.group(1) if m else path.stem[:8]


def _raw_amount(row: pd.Series) -> str:
    place = str(row.get("资金区间_页面") or "").strip()
    if place:
        return re.sub(r"^Board:\s*", "", place, flags=re.I).strip()

    amt = str(row.get("起存金额_页面") or "").strip()
    if amt:
        return amt

    prod = str(row.get("产品或档位") or "").strip()
    if not prod:
        return ""

    for sep in ("—", "–", " - "):
        if sep in prod:
            tail = prod.split(sep, 1)[1].strip()
            if tail:
                return tail
    return prod


def _normalize_money_token(token: str) -> str:
    t = token.strip().replace("\xa0", " ")
    t = re.sub(r"\s+", "", t)
    prefix = ""
    if t.upper().startswith("S$"):
        prefix = "S$"
        num_part = t[2:]
    elif t.startswith("$"):
        prefix = "S$"
        num_part = t[1:]
    elif re.match(r"^\d", t):
        prefix = "S$"
        num_part = t
    else:
        return t

    digits = num_part.replace(",", "")
    if digits.isdigit():
        return f"{prefix}{int(digits):,}"
    return f"{prefix}{num_part}"


def _expand_shorthand_number(text: str) -> str:
    def repl(m: re.Match[str]) -> str:
        n = float(m.group(1).replace(",", ""))
        unit = m.group(2).lower()
        if unit.startswith("m"):
            val = int(n * 1_000_000)
        elif unit.startswith("k"):
            val = int(n * 1_000)
        else:
            val = int(n)
        return f"{val:,}"

    return re.sub(r"([\d,.]+)\s*(mio|mil|million|k)\b", repl, text, flags=re.I)


def normalize_amount_label(text: str) -> str:
    """将资金区间统一为 ≥、＞、≤、＜ 符号表达，去掉中英文比较词。"""
    s = _expand_shorthand_number(str(text or "").strip())
    if not s:
        return ""

    s = re.sub(r"^Board:\s*", "", s, flags=re.I).strip()
    s = s.replace("\xa0", " ").strip()
    s = re.sub(r"^(?:CNY|RMB)\s*", "", s, flags=re.I)
    s = re.sub(r"CNY\s*", "", s, flags=re.I)
    while re.search(r"\d'\d{3}", s):
        s = re.sub(r"(\d)'(\d{3})", r"\1\2", s)
    s = re.sub(r"SGD\s*", "S$", s, flags=re.I)
    s = re.sub(r"S\$\s+", "S$", s, flags=re.I)
    s = re.sub(r"\$\s+", "S$", s)
    s = re.sub(r"\s+", " ", s)

    if re.fullmatch(r"Tier\s+\d+", s, flags=re.I):
        return f"档位{s.split()[-1]}"

    m = re.fullmatch(r"^<\s*([\d,]+)$", s, flags=re.I)
    if m:
        return f"＜{_normalize_money_token(m.group(1))}"

    m = re.fullmatch(r"^<=\s*([\d,]+)$", s, flags=re.I)
    if m:
        return f"≤{_normalize_money_token(m.group(1))}"

    m = re.fullmatch(r"^First\s+(?:S\$|\$)?([\d,]+)$", s, flags=re.I)
    if m:
        return f"≤{_normalize_money_token(m.group(1))}"

    m = re.fullmatch(r"^Above\s+(?:S\$|\$)?([\d,]+)$", s, flags=re.I)
    if m:
        return f"≥{_normalize_money_token(m.group(1))}"

    m = re.fullmatch(r"^([\d,]+)\s*\(inclusive\)\s*and\s*above$", s, flags=re.I)
    if m:
        return f"≥{_normalize_money_token(m.group(1))}"

    m = re.fullmatch(r"^Below\s*([\d,]+)\s*k\s*\(exclusive\)$", s, flags=re.I)
    if m:
        return f"＜{_normalize_money_token(str(int(m.group(1).replace(',', '')) * 1000))}"

    m = re.search(r"<=\s*(?:S\$|\$)?([\d,]+)", s)
    if m:
        return f"≤{_normalize_money_token(m.group(1))}"

    m = re.search(r"Special Rates\s*\(\s*>\s*(?:S\$|\$)?([\d,]+)\s*\)", s, flags=re.I)
    if m:
        return f"＞{_normalize_money_token(m.group(1))}"

    m = re.search(
        r"^>\s*(?:S\$|\$)?([\d,]+)\s*-\s*(?:S\$|\$)?([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"＞{_normalize_money_token(m.group(1))}且≤{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s*至\s*(?:S\$|\$)?([\d,]+)\s*以下$",
        s,
    )
    if m:
        return f"≥{_normalize_money_token(m.group(1))}且＜{_normalize_money_token(m.group(2))}"

    m = re.search(r"^(?:S\$|\$)?([\d,]+)\s*以下$", s)
    if m:
        return f"＜{_normalize_money_token(m.group(1))}"

    m = re.search(r"^(?:S\$|\$)?([\d,]+)\s*(?:及以上|以上)$", s)
    if m:
        return f"≥{_normalize_money_token(m.group(1))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s*-\s*(?:S\$|\$)?([\d,]+)\s*\(exclusive\)$",
        s,
        flags=re.I,
    )
    if m:
        return f"＞{_normalize_money_token(m.group(1))}且＜{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s*-\s*(?:S\$|\$)?([\d,]+)\s*&\s*above$",
        s,
        flags=re.I,
    )
    if m:
        return f"≥{_normalize_money_token(m.group(1))}"

    m = re.search(r"^(?:S\$|\$)?([\d,]+)\s*&\s*above$", s, flags=re.I)
    if m:
        return f"≥{_normalize_money_token(m.group(1))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s+(?:and above|AND ABOVE)$",
        s,
        flags=re.I,
    )
    if m:
        return f"≥{_normalize_money_token(m.group(1))}"

    m = re.search(r"^(?:BELOW|Below)\s+(?:S\$|\$)?([\d,]+)$", s, flags=re.I)
    if m:
        return f"＜{_normalize_money_token(m.group(1))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s+to\s+<\s*(?:S\$|\$)?([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"≥{_normalize_money_token(m.group(1))}且＜{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s*-\s*<\s*(?:S\$|\$)?([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"≥{_normalize_money_token(m.group(1))}且＜{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s+TO\s+(?:S\$|\$)?([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"{_normalize_money_token(m.group(1))}-{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^S\$\s*([\d,]+)\s+to\s+S\$\s*([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"{_normalize_money_token(m.group(1))}-{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s+to\s+(?:S\$|\$)?([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"{_normalize_money_token(m.group(1))}-{_normalize_money_token(m.group(2))}"

    m = re.search(
        r"^(?:S\$|\$)?([\d,]+)\s*-\s*(?:S\$|\$)?([\d,]+)$",
        s,
        flags=re.I,
    )
    if m:
        return f"{_normalize_money_token(m.group(1))}-{_normalize_money_token(m.group(2))}"

    s = re.sub(
        r"(?:S\$|\$)?([\d,]+)\s*(?:及以上|以上|&\s*above|&Above|and above|AND ABOVE)\b",
        lambda m: f"≥{_normalize_money_token(m.group(1))}",
        s,
        flags=re.I,
    )
    s = re.sub(
        r"(?:BELOW|Below)\s+(?:S\$|\$)?([\d,]+)",
        lambda m: f"＜{_normalize_money_token(m.group(1))}",
        s,
        flags=re.I,
    )
    s = re.sub(
        r"(?:S\$|\$)?([\d,]+)\s*(?:以下)\b",
        lambda m: f"＜{_normalize_money_token(m.group(1))}",
        s,
    )
    s = re.sub(r"\bto\b", "-", s, flags=re.I)
    s = re.sub(r"\s*-\s*", "-", s)
    s = re.sub(r"\(exclusive\)", "", s, flags=re.I)
    return s.strip(" -")


def format_amount(row: pd.Series) -> str:
    return normalize_amount_label(_raw_amount(row))


@dataclass
class ParsedBounds:
    low: Optional[float] = None
    high: Optional[float] = None
    low_inclusive: bool = True
    high_inclusive: bool = True
    is_tier: bool = False
    tier_num: Optional[int] = None
    raw: str = ""
    parseable: bool = False


def _money_to_float(token: str) -> float:
    return float(re.sub(r"[^\d.]", "", token.replace(",", "")))


def parse_amount_bounds(label: str) -> ParsedBounds:
    s = str(label or "").strip()
    pb = ParsedBounds(raw=s)
    if not s:
        return pb

    m = re.fullmatch(r"档位(\d+)", s)
    if m:
        pb.is_tier = True
        pb.tier_num = int(m.group(1))
        pb.parseable = True
        return pb

    patterns: list[tuple[str, bool]] = [
        (r"＜(?:S\$|\$)?([\d,]+)", False),
        (r"≤(?:S\$|\$)?([\d,]+)", True),
        (r"≥(?:S\$|\$)?([\d,]+)", True),
        (r"＞(?:S\$|\$)?([\d,]+)", False),
        (r"(?:S\$|\$)?([\d,]+)-(?:S\$|\$)?([\d,]+)", True),
        (r"＞(?:S\$|\$)?([\d,]+)且＜(?:S\$|\$)?([\d,]+)", False),
        (r"≥(?:S\$|\$)?([\d,]+)且＜(?:S\$|\$)?([\d,]+)", True),
        (r"＞(?:S\$|\$)?([\d,]+)且≤(?:S\$|\$)?([\d,]+)", False),
    ]

    if m := re.fullmatch(patterns[0][0], s):
        pb.high = _money_to_float(m.group(1))
        pb.high_inclusive = patterns[0][1]
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[1][0], s):
        pb.high = _money_to_float(m.group(1))
        pb.high_inclusive = patterns[1][1]
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[2][0], s):
        pb.low = _money_to_float(m.group(1))
        pb.low_inclusive = patterns[2][1]
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[3][0], s):
        pb.low = _money_to_float(m.group(1))
        pb.low_inclusive = patterns[3][1]
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[4][0], s):
        pb.low = _money_to_float(m.group(1))
        pb.high = _money_to_float(m.group(2))
        pb.low_inclusive = True
        pb.high_inclusive = True
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[5][0], s):
        pb.low = _money_to_float(m.group(1))
        pb.high = _money_to_float(m.group(2))
        pb.low_inclusive = False
        pb.high_inclusive = False
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[6][0], s):
        pb.low = _money_to_float(m.group(1))
        pb.high = _money_to_float(m.group(2))
        pb.low_inclusive = True
        pb.high_inclusive = False
        pb.parseable = True
        return pb
    if m := re.fullmatch(patterns[7][0], s):
        pb.low = _money_to_float(m.group(1))
        pb.high = _money_to_float(m.group(2))
        pb.low_inclusive = False
        pb.high_inclusive = True
        pb.parseable = True
        return pb

    return pb


def _format_bound_value(value: float) -> str:
    return _normalize_money_token(str(int(value)))


def format_merged_bounds_compact(parts: list[ParsedBounds]) -> str:
    """同利率多档合并：有下界写 ≥最低门槛，仅上界写 ＜最高门槛。"""
    numeric = [p for p in parts if p.parseable and not p.is_tier]
    if not numeric:
        return parts[0].raw if parts else ""

    lows = [p.low for p in numeric if p.low is not None]
    highs = [p.high for p in numeric if p.high is not None]
    has_open_low = any(p.low is None and p.high is not None for p in numeric)

    if lows and not has_open_low:
        min_low = min(lows)
        strict = any(
            p.low == min_low and not p.low_inclusive
            for p in numeric
            if p.low is not None
        )
        op = "＞" if strict else "≥"
        return f"{op}{_format_bound_value(min_low)}"

    if highs:
        max_high = max(highs)
        strict = any(
            p.high == max_high and not p.high_inclusive
            for p in numeric
            if p.high is not None
        )
        op = "＜" if strict else "≤"
        return f"{op}{_format_bound_value(max_high)}"

    return parts[0].raw


def format_merged_bounds(low: Optional[float], high: Optional[float], parts: list[ParsedBounds]) -> str:
    return format_merged_bounds_compact(parts)


def merge_amount_labels(labels: list[str]) -> str:
    expanded: list[str] = []
    for label in labels:
        text = str(label or "").strip()
        if not text:
            continue
        parts = [p.strip() for p in re.split(r"\s*/\s*", text) if p.strip()]
        expanded.extend(parts if len(parts) > 1 else [text])

    uniq = list(dict.fromkeys(expanded))
    if len(uniq) == 1:
        return uniq[0]

    parsed = [parse_amount_bounds(label) for label in uniq]
    tier_parts = [p for p in parsed if p.is_tier]
    if len(tier_parts) == len(parsed):
        nums = sorted(p.tier_num for p in tier_parts if p.tier_num is not None)
        if nums:
            return f"档位{nums[0]}" if nums[0] == nums[-1] else f"档位{nums[0]}-档位{nums[-1]}"

    numeric = [p for p in parsed if p.parseable and not p.is_tier]
    if not numeric:
        return uniq[0]

    return format_merged_bounds_compact(numeric)


def _amount_sort_key(text: str) -> float:
    s = str(text or "").strip()
    if s.startswith("档位"):
        try:
            return float(s.replace("档位", ""))
        except ValueError:
            return float("inf")

    if s.startswith(("＜", "<")):
        m = re.search(r"[\d,]+", s[1:])
        if m:
            return float(m.group(0).replace(",", "")) - 0.5

    m = re.search(r"≥(?:S\$|\$)?([\d,]+)", s)
    if m:
        return float(m.group(1).replace(",", ""))

    m = re.search(r"^(?:S\$|\$)?([\d,]+)", s)
    if m:
        return float(m.group(1).replace(",", ""))

    nums = re.findall(r"[\d,]+", s.replace(",", ""))
    if nums:
        try:
            return float(nums[0].replace(",", ""))
        except ValueError:
            pass
    return float("inf")


def rows_for_tenor(df: pd.DataFrame, tenor: str) -> pd.DataFrame:
    """展示各资金区间利率；同区间多利率取最高，同银行同利率合并区间。

    排序口径（Personal / Highest Rate）：
    - 银行间：最高利率降序；同利率时，该最高利率对应的最低起存额升序
    - 银行内：利率降序；同利率时金额升序（越低越好）
    """
    rate_col = f"{tenor}_pct"
    if rate_col not in df.columns:
        return pd.DataFrame(columns=["Bank", "Rate", "Amount"])

    out = df.copy()
    out[rate_col] = pd.to_numeric(out[rate_col], errors="coerce")
    out = out[out[rate_col].notna()].copy()
    if out.empty:
        return pd.DataFrame(columns=["Bank", "Rate", "Amount"])

    out["_bank"] = out["数据来源"].map(display_bank)
    out["_amount"] = out.apply(format_amount, axis=1)
    out = out[out["_amount"].astype(str).str.strip().ne("")].copy()
    if out.empty:
        return pd.DataFrame(columns=["Bank", "Rate", "Amount"])

    deduped = (
        out.groupby(["_bank", "_amount"], as_index=False)[rate_col]
        .max()
        .rename(columns={rate_col: "_rate"})
    )
    deduped["_rate_key"] = deduped["_rate"].round(6)

    collapsed_rows: list[dict[str, object]] = []
    for (bank, _rate_key), grp in deduped.groupby(["_bank", "_rate_key"], sort=False):
        amounts = grp["_amount"].astype(str).tolist()
        merged_amount = merge_amount_labels(amounts)
        collapsed_rows.append(
            {
                "_bank": bank,
                "_rate": float(grp["_rate"].iloc[0]),
                "_amount": merged_amount,
                "_sort_amt": _amount_sort_key(merged_amount),
            }
        )

    collapsed = pd.DataFrame(collapsed_rows)
    bank_rank_rows: list[dict[str, object]] = []
    for bank, grp in collapsed.groupby("_bank", sort=False):
        max_rate = float(grp["_rate"].max())
        min_amt_at_max = float(grp.loc[grp["_rate"] >= max_rate - 1e-9, "_sort_amt"].min())
        bank_rank_rows.append(
            {"_bank": bank, "_max_rate": max_rate, "_min_amt_at_max": min_amt_at_max}
        )
    bank_order = (
        pd.DataFrame(bank_rank_rows)
        .sort_values(by=["_max_rate", "_min_amt_at_max"], ascending=[False, True])["_bank"]
        .tolist()
    )

    rows: list[dict[str, object]] = []
    for bank in bank_order:
        sub = collapsed[collapsed["_bank"] == bank].copy()
        sub = sub.sort_values(by=["_rate", "_sort_amt"], ascending=[False, True])
        for _, r in sub.iterrows():
            rows.append(
                {
                    "Bank": bank,
                    "Rate": float(r["_rate"]) / 100.0,
                    "Amount": r["_amount"],
                }
            )
    return pd.DataFrame(rows)


def _fill_for_tenor(tenor: str) -> PatternFill:
    return PatternFill("solid", fgColor=TENOR_FILL_HEX.get(tenor, "FFFFFF"))


def write_headers(ws) -> None:
    for tenor in TENORS:
        base = TENOR_START_COL[tenor]
        fill = _fill_for_tenor(tenor)
        for offset, label in enumerate((tenor, "Highest Rate", "Amount")):
            cell = ws.cell(HEADER_ROW, base + offset, label)
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = CELL_BORDER
            cell.fill = fill


def _merge_bank_cells(ws, base_col: int, start_row: int, end_row: int, fill: PatternFill) -> None:
    if end_row <= start_row:
        return
    ws.merge_cells(start_row=start_row, start_column=base_col, end_row=end_row, end_column=base_col)
    cell = ws.cell(start_row, base_col)
    cell.alignment = CENTER
    cell.border = CELL_BORDER
    cell.fill = fill


def write_tenor_block(ws, tenor: str, block_df: pd.DataFrame) -> int:
    base = TENOR_START_COL[tenor]
    fill = _fill_for_tenor(tenor)
    if block_df.empty:
        return 0

    banks = block_df["Bank"].tolist()
    run_start = 0
    for i in range(len(banks) + 1):
        if i < len(banks) and banks[i] == banks[run_start]:
            continue
        run_end = i - 1
        for j in range(run_start, run_end + 1):
            r = DATA_START_ROW + j
            bank_cell = ws.cell(r, base, banks[run_start] if j == run_start else None)
            rate_cell = ws.cell(r, base + 1, block_df.iloc[j]["Rate"])
            amt_cell = ws.cell(r, base + 2, block_df.iloc[j]["Amount"])

            rate_cell.number_format = "0.00%"
            for cell in (bank_cell, rate_cell, amt_cell):
                cell.alignment = CENTER
                cell.border = CELL_BORDER
                cell.fill = fill

        start_row = DATA_START_ROW + run_start
        end_row = DATA_START_ROW + run_end
        _merge_bank_cells(ws, base, start_row, end_row, fill)
        run_start = i

    return len(block_df)


def autosize_columns(ws, max_col: int) -> None:
    for col in range(1, max_col + 1):
        letter = ws.cell(1, col).column_letter
        max_len = 10
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row, col).value
            if v is None:
                continue
            max_len = max(max_len, min(len(str(v)), 48))
        ws.column_dimensions[letter].width = max_len + 2


def populate_sgd_board_sheet(ws, market_path: Path) -> dict[str, int]:
    df = pd.read_excel(market_path, sheet_name="新元挂牌利率")
    date_tag = date_tag_from_path(market_path)

    ws.title = "SGD Board Rate"
    ws["A1"] = date_tag
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = CENTER

    write_headers(ws)

    counts: dict[str, int] = {}
    for tenor in TENORS:
        block = rows_for_tenor(df, tenor)
        n = write_tenor_block(ws, tenor, block)
        counts[tenor] = n

    max_col = TENOR_START_COL[TENORS[-1]] + 2
    autosize_columns(ws, max_col)
    return counts


def build_sgd_board_sheet(market_path: Path, output_path: Path) -> dict[str, int]:
    wb = Workbook()
    ws = wb.active
    counts = populate_sgd_board_sheet(ws, market_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return counts


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="从 MarketRateData 生成 SGD Board Rate 表")
    ap.add_argument(
        "--source",
        required=True,
        help="MarketRateData Excel 路径，例如 runs/20260709/MarketRateData_20260709_15.03.xlsx",
    )
    ap.add_argument(
        "--out",
        default=None,
        help="输出 xlsx；默认与 source 同目录下的 RainbowTable_SGD_Board.xlsx",
    )
    return ap.parse_args()


def main() -> None:
    args = parse_args()
    source = Path(args.source).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)

    if args.out:
        out = Path(args.out).resolve()
    else:
        out = source.parent / "RainbowTable_SGD_Board.xlsx"

    counts = build_sgd_board_sheet(source, out)
    print(f"Generated: {out}")
    print(f"Source:    {source}")
    print("Rows:", ", ".join(f"{k}={v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
