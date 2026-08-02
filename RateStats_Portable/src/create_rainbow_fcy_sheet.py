#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 MarketRateData 生成「USD Rate + Other Currency Rates」工作表。

分区（自上而下）：
1. USD Promo Rate — 期限 1M/3M/6M/9M/12M，每组 3 列（期限 / 最高报价 / 金额要求），组间空 1 列
2. USD Board Rates — 期限 7D/1M/3M/6M/9M/12M，同上
3. RMB Board Rate — 宽表：Bank + 各期限利率 + Minimum Amount + Remarks
4. Other Currency Rates — AUD/NZD/EUR/GBP 各 1M/3M/6M + Minimum Amount + Remarks
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from create_rainbow_sgd_board_sheet import (
    _amount_sort_key,
    _raw_amount,
    date_tag_from_path,
    display_bank,
    merge_amount_labels,
    normalize_amount_label,
    rows_for_tenor,
)

USD_PROMO_TENORS = ["1M", "3M", "6M", "9M", "12M"]
USD_BOARD_TENORS = ["7D", "1M", "3M", "6M", "9M", "12M"]
RMB_TENORS = ["1M", "3M", "6M", "9M", "12M", "18M", "24M"]
OTHER_CURRENCIES = ["AUD", "NZD", "EUR", "GBP"]
OTHER_TENORS = ["1M", "3M", "6M"]

TENOR_FILL_HEX = {
    "7D": "E3F2FD",
    "1M": "E8F4FD",
    "3M": "E8F5E9",
    "6M": "FFF8E1",
    "9M": "FCE4EC",
    "12M": "F3E5F5",
    "18M": "E0F2F1",
    "24M": "FFF3E0",
}

PROMO_EXCLUDE_IN_SOURCE = ("board", "挂牌", "fcy td board", "定存挂牌")

thin = Side(border_style="thin", color="000000")
CELL_BORDER = Border(top=thin, left=thin, right=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_FONT = Font(bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFFF00")
TITLE_FONT = Font(bold=True, color="FF0000")
TITLE_ALIGN = Alignment(horizontal="left", vertical="center")
BOC_FONT = Font(color="FF0000")
DATE_FONT = Font(bold=True)


def block_start_col(idx: int) -> int:
    return 1 + idx * 4


def _fill_for_tenor(tenor: str) -> PatternFill:
    return PatternFill("solid", fgColor=TENOR_FILL_HEX.get(tenor, "FFFFFF"))


def _is_boc(bank: object) -> bool:
    return "BOC" in str(bank or "").upper()


def _apply_cell_style(cell, bank: object = None, *, header: bool = False, fill: PatternFill | None = None) -> None:
    cell.alignment = CENTER
    cell.border = CELL_BORDER
    if header:
        cell.font = HEADER_FONT
    elif _is_boc(bank):
        cell.font = BOC_FONT
    if fill is not None:
        cell.fill = fill


def load_fx_board(market_path: Path, currency: str) -> pd.DataFrame:
    return pd.read_excel(market_path, sheet_name=f"挂牌_{currency}", skiprows=1)


def _is_fcy_promo_row(row: pd.Series) -> bool:
    src = str(row.get("数据来源") or "").lower()
    return not any(p in src for p in PROMO_EXCLUDE_IN_SOURCE)


def _currency_matches(cell: object, currency: str) -> bool:
    """币种列可能是纯代码（USD）或带档位（USD (USD5K...)）。"""
    cur = str(cell or "").strip().upper()
    code = str(currency or "").strip().upper()
    if not cur or not code:
        return False
    if cur == code:
        return True
    return cur.startswith(code + " ") or cur.startswith(code + "(")


def _detect_online(*texts: object) -> bool:
    blob = " ".join(str(t or "") for t in texts).lower()
    return any(k in blob for k in ("online", "e-banking", "ebanking", "网上", "网银", "手机银行"))


def compact_usd_amount(text: str) -> str:
    s = str(text or "").strip()
    if not s:
        return ""
    nums = re.findall(r"[\d,]+", s.replace(",", ""))
    if not nums:
        return normalize_amount_label(s).replace("S$", "$")
    n = int(nums[0].replace(",", ""))
    if n >= 1_000_000 and n % 1_000_000 == 0:
        return f"${n // 1_000_000}M"
    if n >= 1_000 and n % 1_000 == 0:
        return f"${n // 1_000}K"
    return f"${n:,}"


def format_promo_amount(row: pd.Series) -> str:
    amt = str(row.get("起存金额_页面") or "").strip()
    src = str(row.get("数据来源") or "")
    note = str(row.get("页面_1M原文") or "")
    cur = str(row.get("币种") or "")
    online = _detect_online(amt, src, note)

    # 工行等：币种/来源里带 USD5K 档位，优先用它，避免从长说明里误取 “1.”
    tier_blob = f"{cur} {src}"
    m_tier = re.search(r"(?:USD|US\$)\s*([\d,.]+)\s*K", tier_blob, re.I)
    m_below = re.search(r"Below\s*(?:USD|US\$)\s*([\d,.]+)\s*K", tier_blob, re.I)

    if "minimum of 10,000 units" in amt.lower() or "10,000 units" in amt.lower():
        base = "$10K"
    elif m_below:
        n = float(m_below.group(1).replace(",", ""))
        base = f"＜${int(n)}K"
    elif m_tier:
        n = float(m_tier.group(1).replace(",", ""))
        base = f"≥${int(n)}K"
    elif amt and not re.match(r"^\s*\d+\.\s*[A-Z]", amt):
        base = compact_usd_amount(amt)
    else:
        nums = re.findall(r"(?:USD|US\$)\s*([\d,]+)", f"{amt} {src} {note}", re.I)
        if not nums:
            nums = re.findall(r"[\d,]+", f"{src} {note}")
        base = compact_usd_amount(nums[0]) if nums else ""

    if online and base:
        return f"{base} online"
    if online:
        return "online"
    return base


def promo_display_bank(source: object) -> str:
    s = str(source or "").strip()
    rules = [
        (r"(?i)\bBOC\b", "BOC"),
        (r"(?i)\bICBC\b", "ICBC"),
        (r"(?i)\bCIMB\b", "CIMB"),
        (r"(?i)Citibank", "CITI"),
        (r"(?i)\bBEA\b", "BEA"),
        (r"(?i)\bRHB\b", "RHB"),
        (r"(?i)\bSBI\b", "SBI"),
        (r"(?i)\bHL Bank\b", "HL Bank"),
        (r"(?i)\bHSBC\b", "HSBC"),
        (r"(?i)\bMaybank\b", "Maybank"),
        (r"(?i)\bOCBC\b", "OCBC"),
        (r"(?i)\bUOB\b", "UOB"),
        (r"(?i)\bSCB\b", "SCB"),
        (r"(?i)\bDBS\b", "DBS"),
    ]
    for pat, name in rules:
        if re.search(pat, s):
            return name
    return display_bank(s)


def usd_board_amount(amount: str) -> str:
    return str(amount or "").replace("S$", "$")


def preprocess_fcy_amount_text(text: str) -> str:
    s = str(text or "").strip()
    if not s:
        return s
    m = re.search(r"Below\s*([\d,]+)\s*k\s*\(exclusive\)", s, flags=re.I)
    if m:
        return f"Below ${int(m.group(1).replace(',', '')) * 1000:,}"
    m = re.search(r"([\d,]+)\s*k\s*\(inclusive\)\s*&\s*above", s, flags=re.I)
    if m:
        return f"${int(m.group(1).replace(',', '')) * 1000:,} and above"
    m = re.fullmatch(r"\$([\d,]+)k", s, flags=re.I)
    if m:
        return f"${int(m.group(1).replace(',', '')) * 1000:,}"
    return s


def format_fcy_amount(row: pd.Series) -> str:
    raw = preprocess_fcy_amount_text(_raw_amount(row))
    label = normalize_amount_label(raw).replace("S$", "$")
    if re.fullmatch(r"\$[\d,]+$", label):
        return f"≥{label}"
    return label


def rows_for_fcy_promo(df: pd.DataFrame, currency: str, tenor: str) -> pd.DataFrame:
    rate_col = f"{tenor}_pct"
    if rate_col not in df.columns:
        return pd.DataFrame(columns=["Bank", "Rate", "Amount"])

    cur_mask = df["币种"].map(lambda v: _currency_matches(v, currency))
    sub = df[cur_mask & df.apply(_is_fcy_promo_row, axis=1)].copy()
    sub[rate_col] = pd.to_numeric(sub[rate_col], errors="coerce")
    sub = sub[sub[rate_col].notna()].copy()
    if sub.empty:
        return pd.DataFrame(columns=["Bank", "Rate", "Amount"])

    sub["_bank"] = sub["数据来源"].map(promo_display_bank)
    rows: list[dict[str, object]] = []
    for bank, grp in sub.groupby("_bank", sort=False):
        pick = grp.sort_values(by=[rate_col, "起存金额_页面"], ascending=[False, True]).iloc[0]
        amount = format_promo_amount(pick)
        rows.append(
            {
                "Bank": bank,
                "Rate": float(pick[rate_col]) / 100.0,
                "Amount": amount,
                "_sort_amt": _amount_sort_key(amount),
            }
        )

    out = (
        pd.DataFrame(rows)
        .sort_values(by=["Rate", "_sort_amt"], ascending=[False, True])
        .drop(columns="_sort_amt")
        .reset_index(drop=True)
    )
    return out


def _merge_bank_cells(ws, base_col: int, start_row: int, end_row: int, fill: PatternFill) -> None:
    if end_row <= start_row:
        return
    ws.merge_cells(start_row=start_row, start_column=base_col, end_row=end_row, end_column=base_col)
    cell = ws.cell(start_row, base_col)
    _apply_cell_style(cell, fill=fill)


def write_side_by_side_section(
    ws,
    header_row: int,
    tenors: list[str],
    blocks: dict[str, pd.DataFrame],
    rate_header: str,
    amount_header: str,
) -> int:
    max_rows = 0
    for idx, tenor in enumerate(tenors):
        base = block_start_col(idx)
        fill = _fill_for_tenor(tenor)
        labels = (tenor, rate_header, amount_header)
        for off, label in enumerate(labels):
            cell = ws.cell(header_row, base + off, label)
            _apply_cell_style(cell, header=True, fill=fill)

        block = blocks.get(tenor, pd.DataFrame(columns=["Bank", "Rate", "Amount"]))
        if block.empty:
            continue

        banks = block["Bank"].tolist()
        run_start = 0
        data_start = header_row + 1
        for i in range(len(banks) + 1):
            if i < len(banks) and banks[i] == banks[run_start]:
                continue
            run_end = i - 1
            for j in range(run_start, run_end + 1):
                r = data_start + j
                bank = banks[run_start]
                bank_cell = ws.cell(r, base, bank if j == run_start else None)
                rate_cell = ws.cell(r, base + 1, block.iloc[j]["Rate"])
                amt_cell = ws.cell(r, base + 2, block.iloc[j]["Amount"])
                rate_cell.number_format = "0.00%"
                for cell in (bank_cell, rate_cell, amt_cell):
                    _apply_cell_style(cell, bank, fill=fill)

            _merge_bank_cells(ws, base, data_start + run_start, data_start + run_end, fill)
            run_start = i

        max_rows = max(max_rows, len(block))

    return max_rows


def write_section_title(ws, row: int, title: str, max_col: int) -> None:
    cell = ws.cell(row, 1, title)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = TITLE_ALIGN
    if max_col > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
        for c in range(2, max_col + 1):
            mc = ws.cell(row, c)
            mc.fill = TITLE_FILL
            mc.border = CELL_BORDER


def build_rmb_board_rows(cnh_df: pd.DataFrame) -> pd.DataFrame:
    banks = sorted({display_bank(b) for b in cnh_df["数据来源"].dropna().unique()})
    rows: list[dict[str, object]] = []

    for bank in banks:
        sub = cnh_df[cnh_df["数据来源"].map(display_bank) == bank].copy()
        amounts: list[str] = []
        online = False
        row_data: dict[str, object] = {"Bank": bank}

        for tenor in RMB_TENORS:
            rate_col = f"{tenor}_pct"
            if rate_col not in sub.columns:
                row_data[tenor] = None
                continue
            sub[rate_col] = pd.to_numeric(sub[rate_col], errors="coerce")
            valid = sub[sub[rate_col].notna()]
            if valid.empty:
                row_data[tenor] = None
                continue
            best = valid[rate_col].max()
            row_data[tenor] = float(best) / 100.0 if pd.notna(best) else None

        for _, r in sub.iterrows():
            amt = format_fcy_amount(r)
            if amt:
                amounts.append(amt)
            if _detect_online(r.get("产品或档位"), r.get("资金区间_页面"), r.get("数据来源")):
                online = True
            if str(r.get("数据来源") or "").upper().startswith("CIMB"):
                online = online or True  # CIMB FCY board tiers are online-only per market notes

        row_data["Minimum Amount"] = merge_amount_labels(amounts) if amounts else ""
        row_data["Remarks"] = "网上存储" if online else ""
        rows.append(row_data)

    rows_df = pd.DataFrame(rows)
    if not rows_df.empty and any(rows_df[t].notna().any() for t in RMB_TENORS if t in rows_df):
        def bank_sort_key(b: str) -> float:
            sub = rows_df[rows_df["Bank"] == b]
            vals = [sub[t].iloc[0] for t in RMB_TENORS if t in sub.columns and pd.notna(sub[t].iloc[0])]
            return max(vals) if vals else -1.0

        rows_df["_sort"] = rows_df["Bank"].map(bank_sort_key)
        rows_df = rows_df.sort_values("_sort", ascending=False).drop(columns="_sort")
    return rows_df.reset_index(drop=True)


def _best_board_rate(df: pd.DataFrame, bank: str, tenor: str) -> float | None:
    rate_col = f"{tenor}_pct"
    if rate_col not in df.columns:
        return None
    sub = df[df["数据来源"].map(display_bank) == bank].copy()
    if sub.empty:
        # 兼容促销命名（如 BOC SG / Citibank）与挂牌短名混用
        sub = df[df["数据来源"].map(promo_display_bank) == bank].copy()
    sub[rate_col] = pd.to_numeric(sub[rate_col], errors="coerce")
    vals = sub[rate_col].dropna()
    if vals.empty:
        return None
    # 官网偶发用 0.0001 表示近零挂牌；这类值不当作有效报价
    vals = vals[vals >= 0.01]
    if vals.empty:
        return None
    best = vals.max()
    return float(best) / 100.0


def _best_promo_rate(
    promo_df: pd.DataFrame,
    bank: str,
    currency: str,
    tenor: str,
) -> tuple[float | None, pd.Series | None]:
    """返回 (小数利率, 选中的促销行)；无促销则 (None, None)。"""
    rate_col = f"{tenor}_pct"
    if promo_df is None or promo_df.empty or rate_col not in promo_df.columns:
        return None, None
    cur_mask = promo_df["币种"].map(lambda v: _currency_matches(v, currency))
    sub = promo_df[cur_mask & promo_df.apply(_is_fcy_promo_row, axis=1)].copy()
    if sub.empty:
        return None, None
    sub["_bank"] = sub["数据来源"].map(promo_display_bank)
    sub = sub[sub["_bank"] == bank].copy()
    if sub.empty:
        return None, None
    sub[rate_col] = pd.to_numeric(sub[rate_col], errors="coerce")
    sub = sub[sub[rate_col].notna() & (sub[rate_col] >= 0.01)]
    if sub.empty:
        return None, None
    pick = sub.sort_values(by=[rate_col, "起存金额_页面"], ascending=[False, True]).iloc[0]
    return float(pick[rate_col]) / 100.0, pick


def build_other_currency_rows(
    boards: dict[str, pd.DataFrame],
    promo_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """
    Other Currency Rates：同一银行×币种×期限，有促销用促销最高价，否则回退挂牌最高价。
    """
    banks: set[str] = set()
    for df in boards.values():
        if df is None or df.empty or "数据来源" not in df.columns:
            continue
        banks.update(promo_display_bank(b) for b in df["数据来源"].dropna().unique())
        banks.update(display_bank(b) for b in df["数据来源"].dropna().unique())

    if promo_df is not None and not promo_df.empty and "数据来源" in promo_df.columns:
        for cur in OTHER_CURRENCIES:
            cur_mask = promo_df["币种"].map(lambda v: _currency_matches(v, cur))
            sub = promo_df[cur_mask & promo_df.apply(_is_fcy_promo_row, axis=1)]
            banks.update(promo_display_bank(b) for b in sub["数据来源"].dropna().unique())

    banks = {b for b in banks if b}

    rows: list[dict[str, object]] = []
    for bank in sorted(banks):
        row: dict[str, object] = {"Bank": bank}
        amounts: list[str] = []
        online = False
        sort_vals: list[float] = []

        for curr, df in boards.items():
            board_sub = pd.DataFrame()
            if df is not None and not df.empty and "数据来源" in df.columns:
                board_sub = df[
                    (df["数据来源"].map(display_bank) == bank)
                    | (df["数据来源"].map(promo_display_bank) == bank)
                ]

            curr_used_promo = False
            for tenor in OTHER_TENORS:
                col_name = f"{curr}_{tenor}"
                promo_rate, promo_pick = _best_promo_rate(promo_df, bank, curr, tenor)
                if promo_rate is not None:
                    row[col_name] = promo_rate
                    sort_vals.append(promo_rate)
                    curr_used_promo = True
                    if promo_pick is not None:
                        amt = format_promo_amount(promo_pick)
                        if amt:
                            amounts.append(amt)
                        if _detect_online(
                            promo_pick.get("起存金额_页面"),
                            promo_pick.get("数据来源"),
                            promo_pick.get("页面_1M原文"),
                        ):
                            online = True
                    continue

                rate = _best_board_rate(df, bank, tenor) if df is not None else None
                row[col_name] = rate
                if rate is not None:
                    sort_vals.append(rate)

            # 该币种没有促销时，用挂牌档位补 Minimum Amount
            if not curr_used_promo:
                for _, r in board_sub.iterrows():
                    amt = format_fcy_amount(r)
                    if amt:
                        amounts.append(amt)
                    if _detect_online(
                        r.get("产品或档位"), r.get("资金区间_页面"), r.get("数据来源")
                    ):
                        online = True
                    if str(r.get("数据来源") or "").upper().startswith("CIMB"):
                        online = True
            elif str(bank).upper() == "CIMB":
                online = True

        # 没有任何期限利率的银行不输出（避免空行）
        if not sort_vals:
            continue

        row["Minimum Amount"] = merge_amount_labels(amounts) if amounts else ""
        row["Remarks"] = "网上存储" if online else ""
        row["_sort"] = max(sort_vals)
        rows.append(row)

    if not rows:
        return pd.DataFrame(columns=["Bank", "Minimum Amount", "Remarks"])

    out = pd.DataFrame(rows).sort_values("_sort", ascending=False).drop(columns="_sort")
    return out.reset_index(drop=True)


def write_rmb_section(ws, start_row: int, df: pd.DataFrame) -> int:
    headers = ["Bank", *RMB_TENORS, "Minimum Amount", "Remarks"]
    hdr_row = start_row
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(hdr_row, c, h)
        _apply_cell_style(cell, header=True, fill=_fill_for_tenor("1M"))

    for i, (_, row) in enumerate(df.iterrows()):
        r = hdr_row + 1 + i
        bank = row["Bank"]
        ws.cell(r, 1, bank)
        _apply_cell_style(ws.cell(r, 1), bank)
        for j, tenor in enumerate(RMB_TENORS, start=2):
            val = row.get(tenor)
            cell = ws.cell(r, j, val if pd.notna(val) else None)
            if val is not None:
                cell.number_format = "0.00%"
            _apply_cell_style(cell, bank)
        ws.cell(r, len(RMB_TENORS) + 2, row.get("Minimum Amount") or "")
        ws.cell(r, len(RMB_TENORS) + 3, row.get("Remarks") or "")
        _apply_cell_style(ws.cell(r, len(RMB_TENORS) + 2), bank)
        _apply_cell_style(ws.cell(r, len(RMB_TENORS) + 3), bank)

    return len(df)


def write_other_currency_section(ws, start_row: int, df: pd.DataFrame) -> int:
    hdr1 = start_row
    hdr2 = start_row + 1
    col = 1
    ws.cell(hdr1, col, "Bank")
    ws.cell(hdr2, col, "")
    _apply_cell_style(ws.cell(hdr1, col), header=True, fill=_fill_for_tenor("3M"))
    _apply_cell_style(ws.cell(hdr2, col), header=True, fill=_fill_for_tenor("3M"))
    col += 1

    for curr in OTHER_CURRENCIES:
        ws.merge_cells(start_row=hdr1, start_column=col, end_row=hdr1, end_column=col + 2)
        c1 = ws.cell(hdr1, col, curr)
        _apply_cell_style(c1, header=True, fill=_fill_for_tenor("6M"))
        for j, tenor in enumerate(OTHER_TENORS):
            c = ws.cell(hdr2, col + j, tenor)
            _apply_cell_style(c, header=True, fill=_fill_for_tenor(tenor))
        col += 3

    for label in ("Minimum Amount", "Remarks"):
        ws.merge_cells(start_row=hdr1, start_column=col, end_row=hdr2, end_column=col)
        c = ws.cell(hdr1, col, label)
        _apply_cell_style(c, header=True, fill=_fill_for_tenor("9M"))
        col += 1

    data_start = hdr2 + 1
    for i, (_, row) in enumerate(df.iterrows()):
        r = data_start + i
        bank = row["Bank"]
        ws.cell(r, 1, bank)
        _apply_cell_style(ws.cell(r, 1), bank)
        c = 2
        for curr in OTHER_CURRENCIES:
            for tenor in OTHER_TENORS:
                val = row.get(f"{curr}_{tenor}")
                cell = ws.cell(r, c, val if pd.notna(val) else None)
                if val is not None:
                    cell.number_format = "0.00%"
                _apply_cell_style(cell, bank)
                c += 1
        ws.cell(r, c, row.get("Minimum Amount") or "")
        ws.cell(r, c + 1, row.get("Remarks") or "")
        _apply_cell_style(ws.cell(r, c), bank)
        _apply_cell_style(ws.cell(r, c + 1), bank)

    note_row = data_start + len(df) + 1
    ws.cell(note_row, 1, "note: 有促销取促销，无促销再取挂牌；CIMB 外币多为网上存储")
    return len(df) + 2


def autosize_sheet(ws, max_col: int) -> None:
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        max_len = 8
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row, col).value
            if v is None:
                continue
            max_len = max(max_len, min(len(str(v)), 40))
        ws.column_dimensions[letter].width = max_len + 2


def populate_fcy_sheet(ws, market_path: Path) -> None:
    fcy_promo = pd.read_excel(market_path, sheet_name="外币定存促销")
    usd_board = load_fx_board(market_path, "USD")
    cnh_board = load_fx_board(market_path, "CNH")
    other_boards = {c: load_fx_board(market_path, c) for c in OTHER_CURRENCIES}

    date_tag = date_tag_from_path(market_path)
    ws.title = "USD Rate + Other Currency Rates"
    ws["A1"] = date_tag
    ws["A1"].font = DATE_FONT
    ws["A1"].alignment = CENTER

    promo_blocks = {t: rows_for_fcy_promo(fcy_promo, "USD", t) for t in USD_PROMO_TENORS}
    board_blocks = {}
    for t in USD_BOARD_TENORS:
        block = rows_for_tenor(usd_board, t)
        if not block.empty:
            block = block.copy()
            block["Amount"] = block["Amount"].map(usd_board_amount)
        board_blocks[t] = block
    rmb_df = build_rmb_board_rows(cnh_board)
    other_df = build_other_currency_rows(other_boards, fcy_promo)

    side_max_col = block_start_col(max(len(USD_PROMO_TENORS), len(USD_BOARD_TENORS)) - 1) + 2
    row = 2

    write_section_title(ws, row, "USD Promo Rate", side_max_col)
    row += 1
    promo_rows = write_side_by_side_section(
        ws, row, USD_PROMO_TENORS, promo_blocks, "最高报价", "金额要求"
    )
    row += 1 + promo_rows + 1

    write_section_title(ws, row, "USD Board Rates", side_max_col)
    row += 1
    board_rows = write_side_by_side_section(
        ws, row, USD_BOARD_TENORS, board_blocks, "Highest Rate", "Amount"
    )
    row += 1 + board_rows + 1

    rmb_max_col = 2 + len(RMB_TENORS) + 1
    write_section_title(ws, row, "RMB Board Rate", rmb_max_col)
    row += 1
    rmb_n = write_rmb_section(ws, row, rmb_df)
    row += 1 + rmb_n + 1

    other_max_col = 1 + len(OTHER_CURRENCIES) * 3 + 2
    write_section_title(ws, row, "Other Currency Rates", other_max_col)
    row += 1
    write_other_currency_section(ws, row, other_df)

    autosize_sheet(ws, max(side_max_col, rmb_max_col, other_max_col))


def build_fcy_sheet(market_path: Path, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    populate_fcy_sheet(ws, market_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="从 MarketRateData 生成 USD + Other Currency 彩虹表")
    ap.add_argument("--source", required=True, help="MarketRateData Excel 路径")
    ap.add_argument("--out", default=None, help="输出 xlsx 路径")
    return ap.parse_args()


def main() -> None:
    args = parse_args()
    source = Path(args.source).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    out = Path(args.out).resolve() if args.out else source.parent / "RainbowTable_FCY.xlsx"
    build_fcy_sheet(source, out)
    print(f"Generated: {out}")
    print(f"Source:    {source}")


if __name__ == "__main__":
    main()
