#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 MarketRateData 生成完整彩虹表工作簿（三张 sheet）：
- SGD Promotional Rate
- SGD Board Rate
- USD Rate + Other Currency Rates
"""

from __future__ import annotations

import os
import re
from pathlib import Path

from openpyxl import Workbook

from create_rainbow_fcy_sheet import populate_fcy_sheet
from create_rainbow_sgd_board_sheet import populate_sgd_board_sheet
from create_rainbow_sgd_promo_sheet import populate_sgd_promo_sheet


SHEET_ORDER = (
    "SGD Promotional Rate",
    "SGD Board Rate",
    "USD Rate + Other Currency Rates",
)


def rainbow_zh_filename(run_tag: str) -> str:
    return f"彩虹表_按MarketRateData更新_{run_tag}.xlsx"


def run_tag_from_market_path(market_path: Path) -> str:
    m = re.search(r"MarketRateData_(\d{8}_\d{2}\.\d{2})", market_path.name, flags=re.I)
    if m:
        return m.group(1)
    m = re.search(r"(\d{8}_\d{2}\.\d{2})", market_path.stem)
    if m:
        return m.group(1)
    return market_path.stem


def build_rainbow_workbook(market_path: Path, output_path: Path) -> dict[str, object]:
    market_path = market_path.resolve()
    if not market_path.is_file():
        raise FileNotFoundError(market_path)

    wb = Workbook()
    ws_promo = wb.active
    ws_promo.title = SHEET_ORDER[0]
    promo_counts = populate_sgd_promo_sheet(ws_promo, market_path)

    ws_board = wb.create_sheet(SHEET_ORDER[1])
    board_counts = populate_sgd_board_sheet(ws_board, market_path)

    ws_fcy = wb.create_sheet(SHEET_ORDER[2])
    populate_fcy_sheet(ws_fcy, market_path)

    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {
        "output": str(output_path),
        "source": str(market_path),
        "promo_rows": promo_counts,
        "board_rows": board_counts,
    }


def build_rainbow_workbook_for_run(market_path: Path, out_dir: Path, run_tag: str) -> Path:
    out_path = out_dir / rainbow_zh_filename(run_tag)
    build_rainbow_workbook(market_path, out_path)
    return out_path
