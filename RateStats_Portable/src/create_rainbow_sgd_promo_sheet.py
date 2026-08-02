#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 MarketRateData 生成「SGD Promotional Rate」工作表。

版式：
- A1：日期 yyyymmdd
- 第 2 行起每 3 列一组：期限 / Promo Rate / Condition
- 期限顺序：1M, 3M, 6M, 9M, 12M, 18M, 24M（连续列，无空列）
- 每组浅色底色区分；同期限内按利率降序，且同一银行的多档数据排在一起
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from create_rainbow_sgd_board_sheet import _amount_sort_key

TENORS = ["1M", "3M", "6M", "9M", "12M", "18M", "24M"]

# 每期限 3 列，连续排列：1M=A-C, 3M=D-F, 6M=G-I, ...
TENOR_START_COL = {tenor: 1 + idx * 3 for idx, tenor in enumerate(TENORS)}

# 浅色分组底色（与旧彩虹表风格接近）
TENOR_FILL_HEX = {
    "1M": "E8F4FD",
    "3M": "E8F5E9",
    "6M": "FFF8E1",
    "9M": "FCE4EC",
    "12M": "F3E5F5",
    "18M": "E0F2F1",
    "24M": "FFF3E0",
}

HEADER_ROW = 2
DATA_START_ROW = 3

thin = Side(border_style="thin", color="000000")
CELL_BORDER = Border(top=thin, left=thin, right=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_FONT = Font(bold=True)


def _promo_hints(prod: str, place: str) -> str:
    text_blob = f"{prod} {place}".lower()
    hints: list[str] = []
    if any(k in text_blob for k in ("fresh fund", "fresh funds", "新资金", "新人")):
        hints.append("fresh funds")
    if any(k in text_blob for k in ("online", "e-banking", "mobile", "网上", "网银", "手机银行")):
        hints.append("Online")
    return " / ".join(hints)


def _tier_label_from_product(prod: str) -> str:
    """从产品/客群字段提取能区分档位的短标签。"""
    if not prod:
        return ""
    if prod.startswith("For "):
        return prod
    if prod.startswith("新币定存"):
        return re.sub(r"^新币定存\s*[-–—]\s*", "", prod).strip()
    for sep in ("—", "–"):
        if sep in prod:
            head, tail = prod.split(sep, 1)
            head, tail = head.strip(), tail.strip()
            if not tail:
                continue
            low = head.lower()
            if "new-to-bank" in low or "new to bank" in low:
                return f"New-To-Bank — {tail}"
            if "citigold" in low and "new funds" in low:
                return f"Citigold New Funds — {tail}"
            if "new funds" in low:
                return f"New Funds — {tail}"
            if "investment bundle" in low:
                return f"Investment Bundle — {tail}"
            if tail and len(head) <= 48:
                return tail
    m = re.search(r"Minimum Deposit\s*(.+)$", prod, re.I)
    if m:
        return m.group(1).strip()
    if " - " in prod:
        left, right = prod.rsplit(" - ", 1)
        if right.strip() and len(left) <= 48:
            return right.strip()
    return ""


def _looks_like_shared_amount_note(amt: str) -> bool:
    """起存金额字段其实是整页说明、无法区分档位时返回 True。"""
    low = amt.lower()
    if ";" in amt and "counter" in low and ("e-banking" in low or "ebanking" in low):
        return True
    if "minimum deposit amount over counter" in low:
        return True
    return False


def _amount_from_tier_label(tier: str) -> str:
    t = tier.replace("\xa0", " ")
    if re.search(r"20\s*k", t, re.I):
        return "S$20,000"
    if re.search(r"500", t, re.I) and not re.search(r"20\s*k", t, re.I):
        return "S$500"
    m = re.search(r"(?:S\$|US\$|USD)\s*[\d,]+", t, re.I)
    if m:
        return re.sub(r"\s+", "", m.group(0))
    m = re.search(r"\$[\d,]+", t)
    if m:
        return m.group(0)
    return ""


def _deposit_amount_label(amt: str, tier: str) -> str:
    """提取可展示的起存金额（优先页面字段，否则从档位推断）。"""
    if amt and not _looks_like_shared_amount_note(amt):
        if "客户分层" in amt or "非金额分层" in amt:
            return ""
        return amt
    inferred = _amount_from_tier_label(tier)
    if inferred:
        low = tier.lower().replace("\xa0", " ")
        if "20k" in low:
            return f"{inferred} (counter)"
        if "500" in low:
            return f"{inferred} (e-banking)"
        return inferred
    return ""


def _part_already_covered(part: str, existing: list[str]) -> bool:
    p = re.sub(r"\s+", "", part.lower())
    for e in existing:
        e_norm = re.sub(r"\s+", "", e.lower())
        if p == e_norm or p in e_norm or e_norm in p:
            return True
    return False


def format_condition(row: pd.Series) -> str:
    prod = str(row.get("产品或客群") or "").strip()
    amt = str(row.get("起存金额_页面") or "").strip()
    place = str(row.get("资金区间_页面") or "").strip()

    tier = _tier_label_from_product(prod)
    hints = _promo_hints(prod, place)
    parts: list[str] = []

    deposit = _deposit_amount_label(amt, tier)
    # 中行「起存500」与 S$500 同义，有金额字段时不再重复写起存标签
    tier_is_boc_min_only = bool(re.match(r"^起存\s*[\d,]+$", tier or ""))
    if tier and not (tier_is_boc_min_only and deposit):
        parts.append(tier)

    if deposit and not _part_already_covered(deposit, parts):
        parts.append(deposit)

    if not parts:
        if prod:
            parts.append(prod if len(prod) <= 100 else prod[:97] + "...")
        elif amt:
            parts.append(amt)
        elif place and len(place) <= 80:
            parts.append(place)

    if hints and not _part_already_covered(hints, parts):
        parts.append(hints)

    return " / ".join(parts)


def display_bank(name: object) -> str:
    s = str(name or "").strip()
    return re.sub(r"\s+SG$", "", s, flags=re.I).strip() or s


def date_tag_from_path(path: Path) -> str:
    m = re.search(r"MarketRateData_(\d{8})", path.name)
    if m:
        return m.group(1)
    m = re.search(r"(\d{8})", path.stem)
    return m.group(1) if m else path.stem[:8]


def _promo_amount_sort_key(row: pd.Series) -> float:
    """促销行起存金额排序键：越低越好；无法解析时靠后。"""
    prod = str(row.get("产品或客群") or "").strip()
    amt = str(row.get("起存金额_页面") or "").strip()
    tier = _tier_label_from_product(prod)
    deposit = _deposit_amount_label(amt, tier)
    if deposit:
        return _amount_sort_key(deposit)
    if amt and not _looks_like_shared_amount_note(amt):
        return _amount_sort_key(amt)
    if tier:
        return _amount_sort_key(_amount_from_tier_label(tier) or tier)
    return float("inf")


def _boc_mobile_tier_kind(prod: str) -> str:
    """识别中行手机银行促销两档：'500' / '200k'；其它返回空。"""
    p = str(prod or "")
    if "200,000" in p or "200000" in p:
        return "200k"
    if re.search(r"起存\s*500\b", p) or re.search(r"S\$\s*500\b", p):
        return "500"
    return ""


def _collapse_boc_identical_mobile_tiers(sub: pd.DataFrame, rate_col: str) -> pd.DataFrame:
    """彩虹表展示：BOC 起存500/200k 若该期限利率相同，只保留较低起存档。

    Market/解析仍保留两档，便于官网日后对该期限拆出不同利率。
    """
    if sub.empty or str(sub.iloc[0].get("_bank") or "") != "BOC":
        return sub
    if "产品或客群" not in sub.columns:
        return sub

    work = sub.copy()
    work["_boc_tier"] = work["产品或客群"].map(_boc_mobile_tier_kind)
    tiered = work[work["_boc_tier"].isin(["500", "200k"])]
    others = work[~work["_boc_tier"].isin(["500", "200k"])]
    if tiered.empty:
        return sub

    keep_idx: list[object] = []
    for _, grp in tiered.groupby(rate_col, sort=False):
        kinds = set(grp["_boc_tier"].tolist())
        if kinds >= {"500", "200k"}:
            # 同利率两档并存：展示起存更低的 500 档
            keep_idx.extend(grp.loc[grp["_boc_tier"] == "500"].index.tolist()[:1])
        else:
            keep_idx.extend(grp.index.tolist())

    collapsed = work.loc[keep_idx]
    out = pd.concat([collapsed, others], axis=0)
    return out.drop(columns=["_boc_tier"], errors="ignore")


def rows_for_tenor(df: pd.DataFrame, tenor: str) -> pd.DataFrame:
    """同期限全部促销记录。

    排序口径：
    - 银行间：最高利率降序；同利率时，该最高利率对应的最低起存额升序
    - 银行内：利率降序；同利率时起存额升序（越低越好）
    """
    rate_col = f"{tenor}_pct"
    if rate_col not in df.columns:
        return pd.DataFrame(columns=["Bank", "Rate", "Condition"])

    out = df.copy()
    out[rate_col] = pd.to_numeric(out[rate_col], errors="coerce")
    out = out[out[rate_col].notna()].copy()
    if out.empty:
        return pd.DataFrame(columns=["Bank", "Rate", "Condition"])

    out["_bank"] = out["数据来源"].map(display_bank)
    out = out.drop_duplicates(subset=["_bank", rate_col, "起存金额_页面", "产品或客群"], keep="first")
    out["_sort_amt"] = out.apply(_promo_amount_sort_key, axis=1)

    bank_rank_rows: list[dict[str, object]] = []
    for bank, grp in out.groupby("_bank", sort=False):
        max_rate = float(grp[rate_col].max())
        min_amt_at_max = float(grp.loc[grp[rate_col] >= max_rate - 1e-9, "_sort_amt"].min())
        bank_rank_rows.append(
            {"_bank": bank, "_max_rate": max_rate, "_min_amt_at_max": min_amt_at_max}
        )
    bank_order = (
        pd.DataFrame(bank_rank_rows)
        .sort_values(by=["_max_rate", "_min_amt_at_max"], ascending=[False, True])["_bank"]
        .tolist()
    )

    rows: list[dict[str, object]] = []
    for bank in bank_order:
        sub = out[out["_bank"] == bank].sort_values(
            by=[rate_col, "_sort_amt"], ascending=[False, True]
        )
        sub = _collapse_boc_identical_mobile_tiers(sub, rate_col)
        sub = sub.sort_values(by=[rate_col, "_sort_amt"], ascending=[False, True])
        for _, r in sub.iterrows():
            rows.append(
                {
                    "Bank": bank,
                    "Rate": float(r[rate_col]) / 100.0,
                    "Condition": format_condition(r),
                }
            )
    return pd.DataFrame(rows)


def _fill_for_tenor(tenor: str) -> PatternFill:
    return PatternFill("solid", fgColor=TENOR_FILL_HEX.get(tenor, "FFFFFF"))


def write_headers(ws) -> None:
    for tenor in TENORS:
        base = TENOR_START_COL[tenor]
        fill = _fill_for_tenor(tenor)
        for offset, label in enumerate((tenor, "Promo Rate", "Condition")):
            cell = ws.cell(HEADER_ROW, base + offset, label)
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = CELL_BORDER
            cell.fill = fill


def _merge_bank_cells(ws, base_col: int, start_row: int, end_row: int, fill: PatternFill) -> None:
    "同银行多行时合并银行名列，仅保留首个单元格显示名称。"
    if end_row <= start_row:
        return
    ws.merge_cells(start_row=start_row, start_column=base_col, end_row=end_row, end_column=base_col)
    cell = ws.cell(start_row, base_col)
    cell.alignment = CENTER
    cell.border = CELL_BORDER
    cell.fill = fill


def write_tenor_block(ws, tenor: str, block_df: pd.DataFrame) -> int:
    base = TENOR_START_COL[tenor]
    fill = _fill_for_tenor(tenor)
    if block_df.empty:
        return 0

    banks = block_df["Bank"].tolist()
    run_start = 0
    for i in range(len(banks) + 1):
        if i < len(banks) and banks[i] == banks[run_start]:
            continue
        run_end = i - 1
        for j in range(run_start, run_end + 1):
            r = DATA_START_ROW + j
            bank_cell = ws.cell(r, base, banks[run_start] if j == run_start else None)
            rate_cell = ws.cell(r, base + 1, block_df.iloc[j]["Rate"])
            cond_cell = ws.cell(r, base + 2, block_df.iloc[j]["Condition"])

            rate_cell.number_format = "0.00%"
            for cell in (bank_cell, rate_cell, cond_cell):
                cell.alignment = CENTER
                cell.border = CELL_BORDER
                cell.fill = fill

        start_row = DATA_START_ROW + run_start
        end_row = DATA_START_ROW + run_end
        _merge_bank_cells(ws, base, start_row, end_row, fill)
        run_start = i

    return len(block_df)


def autosize_columns(ws, max_col: int) -> None:
    for col in range(1, max_col + 1):
        letter = ws.cell(1, col).column_letter
        max_len = 10
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row, col).value
            if v is None:
                continue
            max_len = max(max_len, min(len(str(v)), 48))
        ws.column_dimensions[letter].width = max_len + 2


def populate_sgd_promo_sheet(ws, market_path: Path) -> dict[str, int]:
    df = pd.read_excel(market_path, sheet_name="新元定存促销")
    date_tag = date_tag_from_path(market_path)

    ws.title = "SGD Promotional Rate"
    ws["A1"] = date_tag
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = CENTER

    write_headers(ws)

    counts: dict[str, int] = {}
    for tenor in TENORS:
        block = rows_for_tenor(df, tenor)
        n = write_tenor_block(ws, tenor, block)
        counts[tenor] = n

    max_col = TENOR_START_COL[TENORS[-1]] + 2
    autosize_columns(ws, max_col)
    return counts


def build_sgd_promo_sheet(market_path: Path, output_path: Path) -> dict[str, int]:
    wb = Workbook()
    ws = wb.active
    counts = populate_sgd_promo_sheet(ws, market_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return counts


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="从 MarketRateData 生成 SGD Promotional Rate 表")
    ap.add_argument(
        "--source",
        required=True,
        help="MarketRateData Excel 路径，例如 runs/20260709/MarketRateData_20260709_15.03.xlsx",
    )
    ap.add_argument(
        "--out",
        default=None,
        help="输出 xlsx；默认与 source 同目录下的 RainbowTable_SGD_Promo.xlsx",
    )
    return ap.parse_args()


def main() -> None:
    args = parse_args()
    source = Path(args.source).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)

    if args.out:
        out = Path(args.out).resolve()
    else:
        out = source.parent / "RainbowTable_SGD_Promo.xlsx"

    counts = build_sgd_promo_sheet(source, out)
    print(f"Generated: {out}")
    print(f"Source:    {source}")
    print("Rows:", ", ".join(f"{k}={v}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
